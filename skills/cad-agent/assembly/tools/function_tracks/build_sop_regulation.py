# -*- coding: utf-8 -*-
"""生成《OLSK Large CNC V3 装配工艺规程 v1》。

把官方 Workbook 的 114 步升级为工艺卡：功能序、零件清单、紧固件清单（含工具与推荐扭矩）、
装后调整/检验点、前置条件、官方 Notes/Remarks 原文、待确认问题；并附紧固件总表、工具总表、
扭矩参考表与台架并行计划。

只读输入（全部相对本案例根目录）：
  装配工艺/sop-olsk-workbook.v1.json      官方 114 步逐字搬运
  装配工艺/manual-bench-groups.v1.json    手册源码内建的 preparing 分组与接线步
  本体/S3-工序实例绑定.v1.json             bench_groups / unit_joins / predecessors / movers
  本体/S7-工序内播放序.v3.json              工序内功能序（实体 occ 序）
  功能属性/tracks/A..H/function.json       功能属性（承载件、安装方式、装后调整、顺序约束）
  功能属性/tracks/A..H/order_constraints.json  工序间顺序约束
  功能属性/inputs/Workbook.csv             官方 SOP 原表（校验用）
  功能属性/inputs/HowTo_Blad1.csv          How-To 原文
  功能属性/inputs/OLSK_README.md           机器概况
  功能属性/inputs/OLSK_Large_CNC_V3_Settings.txt  控制器设置（行程等）

输出（只写 工艺文件/装配工艺规程/）：
  sop.json
  OLSK_Large_CNC_V3_装配工艺规程_v1.xlsx
  OLSK_Large_CNC_V3_装配工艺规程_v1.pdf（经 sop.tex，xelatex 两遍）

运行：
  .venv/bin/python 工艺文件/装配工艺规程/build_sop.py
"""
from __future__ import annotations

import csv
import datetime
import hashlib
import json
import re
import subprocess
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
CASE = HERE.parent.parent
OUT_JSON = HERE / "sop.json"
OUT_XLSX = HERE / "OLSK_Large_CNC_V3_装配工艺规程_v1.xlsx"
OUT_TEX = HERE / "sop.tex"
OUT_PDF = HERE / "OLSK_Large_CNC_V3_装配工艺规程_v1.pdf"
DATE = datetime.date.today().strftime("%Y 年 %m 月 %d 日")
VERSION = "v1"
TRACKS = "ABCDEFGH"

# =====================================================================================
# 0. 通用工具
# =====================================================================================


def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def load(rel: str):
    return json.loads((CASE / rel).read_text(encoding="utf-8"))


def lst(x):
    if x is None:
        return []
    if isinstance(x, list):
        return x
    return [x]


OP_RE = re.compile(r"WB-(\d+(?:\.\d+)?)")
BARE_RE = re.compile(r"(?<![\w.-])(\d{2}(?:\.\d)?)(?![\w])")


def ops_in_code(code: str) -> list[str]:
    """从功能线的 subassembly.code / 约束的 from-to 字段里抽出工序号。

    支持 'WB-64.1 Prepare right window / WB-65.1 Prepare left window'
    与 'WB-69.1 / 69.2 / 69.3 / 69.4 Wiring AC 1-4' 两种写法。
    """
    code = str(code or "")
    found = OP_RE.findall(code)
    out = ["WB-" + f for f in found]
    # 'WB-69.1 / 69.2 / 69.3' 形式：把裸号也收进来（只在已出现过 WB- 前缀时）
    if found:
        for seg in re.split(r"[／/、]", code):
            if OP_RE.search(seg):
                continue
            m = BARE_RE.search(seg.strip())
            if m:
                out.append("WB-" + m.group(1))
    seen, res = set(), []
    for o in out:
        if o not in seen:
            seen.add(o)
            res.append(o)
    return res


def first_op(code: str) -> str | None:
    o = ops_in_code(code)
    return o[0] if o else None


def runs(names: list[str]) -> list[tuple[str, int]]:
    """把 ['A','A','B'] 折叠成 [('A',2),('B',1)]（保持次序）。"""
    out: list[list] = []
    for n in names:
        if out and out[-1][0] == n:
            out[-1][1] += 1
        else:
            out.append([n, 1])
    return [(a, b) for a, b in out]


def fmt_runs(rs: list[tuple[str, int]]) -> str:
    return " → ".join(f"{n}×{q}" if q > 1 else n for n, q in rs)


# =====================================================================================
# 1. 紧固件解码：家族 → 头型 → 工具 → 扭矩
# =====================================================================================

FASTENER_RE = re.compile(
    r"^(B-screw|C-screw|HF-screw|Countersunk Screw|Thermoplastic Screw|Hex Screw|Lock Nut|Washer|Standoff|Set Screw)\b"
)

FAMILY = {
    "B-screw": dict(
        head="内六角半圆头",
        std="ISO 7380（手册前缀 B-screw；BOM 写作 ISO 3780，该号非螺钉标准，疑笔误）",
        drive="内六角",
        std_src="推断",
    ),
    "C-screw": dict(head="内六角圆柱头", std="DIN 912 / ISO 4762", drive="内六角", std_src="推断"),
    "HF-screw": dict(head="六角法兰头", std="标准号无记录（手册自造前缀 HF）", drive="外六角", std_src="待确认"),
    "Hex Screw": dict(head="外六角头", std="DIN 933 / ISO 4017", drive="外六角", std_src="推断"),
    "Countersunk Screw": dict(head="内六角沉头", std="DIN 7991 / ISO 10642", drive="内六角", std_src="推断"),
    "Thermoplastic Screw": dict(head="塑料直接紧固自攻螺钉（头型无记录）", std="无记录", drive="内六角（推断）", std_src="待确认"),
    "Lock Nut": dict(head="尼龙嵌件六角锁紧螺母", std="DIN 985 / ISO 10511", drive="外六角", std_src="推断"),
    "Washer": dict(head="平垫圈", std="DIN 125（Large：DIN 9021）", drive="—", std_src="推断"),
    "Standoff": dict(head="双内螺纹隔离柱（FF）", std="无记录", drive="外六角（推断）", std_src="待确认"),
    "Set Screw": dict(head="内六角紧定螺钉", std="DIN 913 / 916", drive="内六角", std_src="推断"),
}

# 内六角对边（mm）
HEX_BUTTON = {"M3": 2, "M4": 2.5, "M5": 3, "M6": 4, "M8": 5, "M10": 6, "M12": 8}       # ISO 7380
HEX_CAP = {"M3": 2.5, "M4": 3, "M5": 4, "M6": 5, "M8": 6, "M10": 8, "M12": 10}          # DIN 912
HEX_CSK = {"M3": 2, "M4": 2.5, "M5": 3, "M6": 4, "M8": 5, "M10": 6, "M12": 8}           # DIN 7991
HEX_SET = {"M3": 1.5, "M4": 2, "M5": 2.5, "M6": 3, "M8": 4}                             # DIN 913
# 外六角对边（mm）
WAF_NUT = {"M3": 5.5, "M4": 7, "M5": 8, "M6": 10, "M8": 13, "M10": 17, "M12": 19}       # DIN 985
WAF_HEX = {"M6": 10, "M8": 13, "M10": 17, "M12": 19}                                    # DIN 933
WAF_FLANGE = {"M6": 10, "M8": 13, "M10": 15, "M12": 18}                                 # DIN 6921 法兰头

# 官方手册唯一一次显式给出的工具（WB-01.1）
OFFICIAL_TOOLS = {"Allen Key 6", "Wrench 14"}

# ISO 898-1 8.8 级、μ_总=0.14、90% 屈服利用率的常用查表值（N·m）
TORQUE_A = {"M3": 1.3, "M4": 3.0, "M5": 5.9, "M6": 10.1, "M8": 24.6, "M10": 48.1, "M12": 84.0}
K_ALU = 0.70   # 被夹件为 5 mm 铝管壁 / 6 mm 铝连接板：按承压折减
K_RESIN = 0.30  # 被夹件为光固化树脂 3D 打印件：按承压折减
TORQUE_THERMO = {"M5": 2.5}   # 热塑自攻螺钉一次拧入
TORQUE_SET = {"M4": 2.0, "M5": 3.0, "M6": 4.0}  # 顶丝（顶电机轴 D 面，加螺纹胶）
PREVAIL = {"M3": 0.2, "M4": 0.4, "M5": 0.5, "M6": 1.0, "M8": 2.0, "M10": 3.0, "M12": 5.0}  # 锁紧螺母自锁力矩

TORQUE_CLASS = {
    "A": dict(
        name="A 钢制螺纹副（通孔螺栓 + 螺母，或已确认旋入钢制采购件）",
        rule="ISO 898-1 8.8 级、总摩擦系数 μ = 0.14 常用查表值，取 90% 屈服利用率",
        k="1.00",
    ),
    "B": dict(
        name="B 旋入/夹紧铝件（机架方管、连接板、角件、铝安装板）——本机默认类别",
        rule=f"A 类值 × {K_ALU:.2f}（被夹件为 5 mm 铝管壁、6 mm 铝连接板或 CNC 铣削铝板，按铝的承压能力折减）；"
             "本机结构件全部为铝，除“通孔螺栓 + 锁紧螺母”外一律默认取 B 类",
        k=f"{K_ALU:.2f}",
    ),
    "C": dict(
        name="C 夹紧树脂 3D 打印件",
        rule=f"A 类值 × {K_RESIN:.2f}（功能线 H 设计规则：树脂接头扭矩必须低于同规格金属件；建议加大垫圈分散压强）",
        k=f"{K_RESIN:.2f}",
    ),
    "D": dict(name="D 热塑自攻螺钉旋入塑料/复合板", rule="按塑料直接紧固行业常见值给定绝对值；一次成型，不得反复拆装同孔", k="—"),
    "E": dict(name="E 紧定螺钉（顶丝）顶电机轴", rule="按小规格顶丝行业常见值给定绝对值；必须顶在轴的平面（D 面）上并加螺纹胶（HowTo H4）", k="—"),
    "F": dict(name="F 尼龙嵌件锁紧螺母", rule="A 类值 + 该规格自锁力矩（尼龙圈附加力矩）；尼龙侧背向螺钉头，拆下不得重复使用（HowTo H2）", k="1.00 + 自锁力矩"),
    "X": dict(name="X 不适用/待确认", rule="垫圈无扭矩；隔离柱与顶丝无规格记录，扭矩待确认", k="—"),
}

# 光固化树脂 3D 打印件（来自功能属性 H 线 U-H-RESIN 的五类清单，按手册命名匹配）
RESIN_KEYS = [
    "Free Bearing Holder", "Fixed Bearing", "Bearing holder", "Bearing Holder", "Ball Screw Spacer",
    "Motor Holder", "Wire Way", "Wire Cover", "wire Cover", "Chain Holder", "Endstop Plate Holder",
    "Pipe Attacher", "Pipe Holder", "Brush Piston Support", "Piston Support", "Window Spacer",
    "Attacher Block", "Flat Spacer", "Pulley Side Spacer", "Shoulder Cover", "Tool Holder Support",
    "Monitor Frame", "Handle Attachment", "Backplate Box",
]
# 铝型材/铝结构件（机架方管、连接板、角件、盖板、面板）
ALU_KEYS = ["Beam", "Profile", "Bracket", "Connector Plate", "Corner", "Panel", "Plate", "Cover", "Frame"]

SIZE_RE = re.compile(r"\bM(\d{1,2})(?:[-x](\d{1,3}))?")


def parse_fastener(name: str) -> dict | None:
    m = FASTENER_RE.match(name)
    if not m:
        return None
    fam = m.group(1)
    sm = SIZE_RE.search(name)
    thread = f"M{sm.group(1)}" if sm else None
    length = f"{sm.group(2)} mm" if (sm and sm.group(2)) else None
    return dict(family=fam, thread=thread, length=length, name=name, large="Large" in name)


def tool_for(f: dict) -> tuple[str, str]:
    """返回 (工具描述, 依据)。"""
    fam, th = f["family"], f["thread"]
    if fam == "Washer":
        return "无（装配辅件）", "—"
    if fam in ("B-screw",):
        s = HEX_BUTTON.get(th)
        return (f"内六角扳手 {s:g} mm" if s else "内六角扳手（规格待定）"), "ISO 7380 半圆头对边表（推断）"
    if fam == "C-screw":
        s = HEX_CAP.get(th)
        return (f"内六角扳手 {s:g} mm" if s else "内六角扳手（规格待定）"), "DIN 912 圆柱头对边表（推断）"
    if fam == "Countersunk Screw":
        s = HEX_CSK.get(th)
        return (f"内六角扳手 {s:g} mm" if s else "内六角扳手（规格待定）"), "DIN 7991 沉头对边表（推断）"
    if fam == "Thermoplastic Screw":
        s = HEX_BUTTON.get(th)
        return (f"内六角扳手 {s:g} mm（头型无记录）" if s else "内六角扳手（规格待定）"), "假定内六角头（待确认）"
    if fam == "HF-screw":
        s = WAF_FLANGE.get(th)
        return (f"套筒/开口扳手 {s:g} mm" if s else "扳手（规格待定）"), "DIN 6921 六角法兰头对边表（推断，HF 标准号无记录）"
    if fam == "Hex Screw":
        s = WAF_HEX.get(th)
        return (f"套筒/开口扳手 {s:g} mm" if s else "扳手（规格待定）"), "DIN 933 六角头对边表（推断）"
    if fam == "Lock Nut":
        s = WAF_NUT.get(th)
        if th == "M8":
            return "扳手 14 mm（官方）", "官方 Workbook WB-01.1 工具栏（DIN 985 标准对边为 13 mm，官方值大一档，见附录 B）"
        return (f"扳手 {s:g} mm" if s else "扳手（规格待定）"), "DIN 985 对边表（推断）"
    if fam == "Set Screw":
        return "内六角扳手（规格待定，手册未给螺纹规格）", "待确认"
    if fam == "Standoff":
        return "尖嘴钳/小扳手（规格待定）", "待确认（FF 8 mm 只给了长度，未给螺纹）"
    return "待确认", "待确认"


def op_substrates(part_names: list[str]) -> tuple[bool, bool]:
    resin = any(any(k in n for k in RESIN_KEYS) for n in part_names)
    alu = any(any(k in n for k in ALU_KEYS) for n in part_names)
    return resin, alu


def rnd(v: float) -> str:
    return f"{v:.1f}" if v < 10 else f"{v:.0f}"


def torque_for(f: dict, part_names: list[str], locknut_threads: set[str], resin_seed: bool = False) -> dict:
    """按 家族 → 同步锁母 → 被夹材料 判定扭矩类别，返回 dict(class, value, text, basis)。"""
    fam, th = f["family"], f["thread"]
    resin, alu = op_substrates(part_names)
    if fam == "Washer":
        return dict(cls="X", value=None, text="—", why="垫圈本身不承受拧紧力矩")
    if fam == "Standoff":
        return dict(cls="X", value=None, text="待确认", why="Standoff FF 8mm 只给长度未给螺纹规格")
    if fam == "Set Screw":
        v = TORQUE_SET.get(th or "M5", 3.0)
        return dict(cls="E", value=v, text=f"{rnd(v)} N·m + 螺纹胶", why="手册未给规格，按 M5 顶丝取值（推断）；HowTo H4 要求顶在轴的平面上")
    if fam == "Thermoplastic Screw":
        v = TORQUE_THERMO.get(th, 2.5)
        return dict(cls="D", value=v, text=f"{rnd(v)} N·m", why="旋入塑料/复合板，一次成型，按塑料直接紧固常见值")
    if fam == "Lock Nut":
        base = TORQUE_A.get(th)
        if base is None:
            return dict(cls="X", value=None, text="待确认", why="螺纹规格未识别")
        v = base + PREVAIL.get(th, 0.0)
        return dict(cls="F", value=v, text=f"{rnd(v)} N·m", why=f"A 类 {rnd(base)} N·m + 自锁力矩 {PREVAIL.get(th,0):.1f} N·m；尼龙侧背向螺钉头（H2）")
    base = TORQUE_A.get(th)
    if base is None:
        return dict(cls="X", value=None, text="待确认", why="螺纹规格未识别")
    # 通孔螺栓 + 同规格尼龙锁紧螺母 → 螺纹副落在钢制螺母上，可取足值
    if th in locknut_threads:
        return dict(cls="A", value=base, text=f"{rnd(base)} N·m", why="本步有同规格 Lock Nut，判为通孔螺栓 + 钢制锁紧螺母，螺纹副为钢-钢（推断）")
    if fam == "HF-screw":
        v = base * K_ALU
        return dict(cls="B", value=v, text=f"{rnd(v)} N·m", why="型材端面对接：咬合连接板并压紧 5 mm 铝管壁（推断）")
    if resin_seed and resin and th in ("M3", "M4", "M5", "M6"):
        v = base * K_RESIN
        return dict(cls="C", value=v, text=f"{rnd(v)} N·m", why="本步的承载基准件是树脂 3D 打印件，整步取树脂折减值（推断）")
    v = base * K_ALU
    why = ("被夹/旋入件为铝型材、铝板或铝支座（推断）" if alu else
           "本机结构件全部为 CNC 铣削铝型材与铝板，默认按铝件折减；若确认此处旋入的是钢制采购件"
           "（滑块、丝杠螺母法兰、电机端面、轴承座、DIN 导轨），可上调至 A 类值 "
           f"{rnd(base)} N·m（推断）")
    return dict(cls="B", value=v, text=f"{rnd(v)} N·m", why=why)


# =====================================================================================
# 2. 读取输入
# =====================================================================================

SRC_FILES = [
    "装配工艺/sop-olsk-workbook.v1.json",
    "装配工艺/manual-bench-groups.v1.json",
    "本体/S3-工序实例绑定.v1.json",
    "本体/S7-工序内播放序.v3.json",
    "功能属性/inputs/Workbook.csv",
    "功能属性/inputs/HowTo_Blad1.csv",
    "功能属性/inputs/OLSK_README.md",
    "功能属性/inputs/OLSK_Large_CNC_V3_Settings.txt",
] + [f"功能属性/tracks/{t}/{f}" for t in TRACKS for f in ("function.json", "order_constraints.json")]

WB = load("装配工艺/sop-olsk-workbook.v1.json")
BENCH_RAW = load("装配工艺/manual-bench-groups.v1.json")
S3 = load("本体/S3-工序实例绑定.v1.json")
S7 = load("本体/S7-工序内播放序.v3.json")

TRACK_FN = {t: load(f"功能属性/tracks/{t}/function.json") for t in TRACKS}
TRACK_OC = {t: load(f"功能属性/tracks/{t}/order_constraints.json") for t in TRACKS}


def read_howtos() -> dict:
    rows = list(csv.reader(open(CASE / "功能属性/inputs/HowTo_Blad1.csv", encoding="utf-8")))
    out: OrderedDict[str, dict] = OrderedDict()
    cur = None
    for r in rows[1:]:
        if not r:
            continue
        hid = r[0].strip()
        if hid:
            cur = hid
            out[cur] = dict(id=hid, title=r[1].strip(), parts=[], tools=[], remarks=[])
        if cur is None:
            continue
        rec = out[cur]
        if len(r) > 2 and r[2].strip():
            rec["parts"].append(r[2].strip())
        if len(r) > 3 and r[3].strip():
            rec["tools"].append(r[3].strip())
        if len(r) > 4 and r[4].strip():
            rec["remarks"].append(" ".join(r[4].split()))
    return out


HOWTO = read_howtos()

SETTINGS = dict(
    line.split("=", 1) for line in (CASE / "功能属性/inputs/OLSK_Large_CNC_V3_Settings.txt").read_text().splitlines()
    if "=" in line
)

# ---- occ_id → 手册零件名 ---------------------------------------------------------
OCC2NAME: dict[str, str] = {}
for _o in S3["operations"]:
    for _m in _o.get("movers", []):
        oid = _m.get("occ_id")
        if not oid:
            continue
        nm = (_m.get("via_glb_names") or [None])[0] or _m.get("product_name") or _m.get("product")
        OCC2NAME[oid] = nm

S3_BY_OP = {o["op"]: o for o in S3["operations"]}
BENCH_OF_OP = {}
for g in S3["bench_groups"]:
    for op in g["ops"]:
        BENCH_OF_OP[op] = g["group_id"]
JOIN_OF_BENCH: dict[str, str] = {}
for o in S3["operations"]:
    for uj in o.get("unit_joins") or []:
        JOIN_OF_BENCH[uj["unit"]] = o["op"]
BENCH_OPS = {g["group_id"]: g["ops"] for g in S3["bench_groups"]}
WIRING_OPS = set()
for raw in BENCH_RAW.get("wiring_steps", []):
    m = re.match(r"^(\d+)(?:_|$)", raw)
    if m:
        d = m.group(1)
        step = d if len(d) <= 2 else f"{d[:2]}.{d[2:]}"
        WIRING_OPS.add("WB-" + step)

# ---- 功能线：subassembly / unit 按工序索引 ---------------------------------------
SUB_BY_OP: dict[str, list[dict]] = defaultdict(list)
UNIT_BY_OP: dict[str, list[dict]] = defaultdict(list)      # 经部套 code 命中的单元（用于设计规则）
UNIT_PRIMARY_BY_OP: dict[str, list[dict]] = defaultdict(list)  # 部套 code 的首个工序号 == 本步（用于待确认问题）
for _t in TRACKS:
    _d = TRACK_FN[_t]
    for _u in _d.get("units", []):
        _u = dict(_u, _track=_t, _track_title=_d.get("title"))
        for _s in _u.get("subassemblies", []):
            _ops = ops_in_code(_s.get("code", ""))
            for _op in _ops:
                SUB_BY_OP[_op].append(dict(_s, _track=_t, _unit=_u["unit_id"], _unit_name=_u.get("name")))
                if all(x["unit_id"] != _u["unit_id"] for x in UNIT_BY_OP[_op]):
                    UNIT_BY_OP[_op].append(_u)
            if _ops and all(x["unit_id"] != _u["unit_id"] for x in UNIT_PRIMARY_BY_OP[_ops[0]]):
                UNIT_PRIMARY_BY_OP[_ops[0]].append(_u)

# 单元级 open_questions / data_gaps 只挂在该单元最早的一步上（避免逐步重复）
UNIT_ANCHOR: dict[str, list[dict]] = defaultdict(list)
_seen_unit: set[tuple] = set()
for _op in [o["op"] for o in WB["operations"]]:
    for _u in UNIT_PRIMARY_BY_OP.get(_op, []):
        _key = (_u["_track"], _u["unit_id"])
        if _key in _seen_unit:
            continue
        _seen_unit.add(_key)
        UNIT_ANCHOR[_op].append(_u)

# ---- 工序间顺序约束 --------------------------------------------------------------
PRE_KINDS = ("MUST_BEFORE", "ACCESS", "WIRE_BEFORE_CLOSE")
HINT_KINDS = ("SAFETY", "SAME_BENCH", "MUST_AFTER")

CONSTRAINTS: list[dict] = []
for t in TRACKS:
    for c in TRACK_OC[t]:
        fr, to = first_op(c.get("from")), first_op(c.get("to"))
        if not fr or not to:
            continue
        CONSTRAINTS.append(dict(
            track=t, kind=c.get("kind"), frm=fr, to=to,
            intra=(fr == to),
            why=" ".join(str(c.get("why") or "").split()),
            source=c.get("source"), confidence=c.get("confidence"),
            parts=lst(c.get("parts")), from_part=c.get("from_part"), to_part=c.get("to_part"),
            status=c.get("status"),
        ))

PRE_BY_OP: dict[str, list[dict]] = defaultdict(list)
HINT_BY_OP: dict[str, list[dict]] = defaultdict(list)
ADJ_BY_OP: dict[str, list[dict]] = defaultdict(list)
INTRA_BY_OP: dict[str, list[dict]] = defaultdict(list)
for c in CONSTRAINTS:
    if c["intra"]:
        INTRA_BY_OP[c["frm"]].append(c)
        continue
    if c["kind"] in PRE_KINDS:
        PRE_BY_OP[c["to"]].append(c)
    elif c["kind"] == "ADJUST_AFTER":
        ADJ_BY_OP[c["frm"]].append(c)
    elif c["kind"] in HINT_KINDS:
        HINT_BY_OP[c["to"]].append(c)
        HINT_BY_OP[c["frm"]].append(c)

# =====================================================================================
# 3. 组装工艺卡
# =====================================================================================

OPS_ORDER = [o["op"] for o in WB["operations"]]
OP_INDEX = {op: i for i, op in enumerate(OPS_ORDER)}

CARDS: list[dict] = []
FAST_TOTAL: Counter = Counter()
FAST_BY_OP: dict[str, list[dict]] = {}
TOOL_TOTAL: dict[str, dict] = {}


def add_tool(name: str, basis: str, op: str):
    rec = TOOL_TOTAL.setdefault(name, dict(name=name, basis=basis, ops=[]))
    if op not in rec["ops"]:
        rec["ops"].append(op)


for wbop in WB["operations"]:
    op = wbop["op"]
    s3 = S3_BY_OP.get(op, {})
    bench = BENCH_OF_OP.get(op)
    join_op = JOIN_OF_BENCH.get(bench) if bench else None
    joins_here = [uj["unit"] for uj in (s3.get("unit_joins") or [])]

    if bench:
        role = f"台架预装组 {bench}（在 {join_op} 整体上机）" if join_op else f"台架预装组 {bench}"
    elif joins_here:
        role = "上机结合步（接收 " + "、".join(joins_here) + " 台架完成件）"
    elif op in WIRING_OPS:
        role = "接线/管路步（手册 wiring 分组）"
    else:
        role = "主线装配步"

    part_names = [p["name"] for p in wbop["parts"]]
    locknut_threads = {parse_fastener(n)["thread"] for n in part_names
                       if (parse_fastener(n) or {}).get("family") == "Lock Nut"}
    locknut_threads = {t for t in locknut_threads if t}
    # 基材判定还要看功能线给出的关键件与本步承载基准件（很多步的承载件是上一步装好的板/梁）
    # 只取"只对应本步一个工序号"的部套的关键件，避免 H 线跨步 code 的关键件串入
    mat_names = list(part_names)
    for _s in SUB_BY_OP.get(op, []):
        if ops_in_code(_s.get("code", "")) == [op]:
            mat_names += [str(x) for x in lst(_s.get("key_parts"))]
    _s7pre = S7["ops"].get(op)
    if _s7pre:
        mat_names += [OCC2NAME.get(i["occ"], "") for i in _s7pre["order"]]
        _seed0 = (_s7pre.get("functional_sort") or {}).get("seed")
        if _seed0:
            mat_names.append(OCC2NAME.get(_seed0, ""))

    _seed_name = OCC2NAME.get((( S7["ops"].get(op) or {}).get("functional_sort") or {}).get("seed"), "") or ""
    resin_seed = any(k in _seed_name for k in RESIN_KEYS)
    resin_parts_here = sorted({n for n in mat_names if n and any(k in n for k in RESIN_KEYS)})

    # --- 零件清单 / 紧固件清单 ---
    parts, fasteners = [], []
    for p in wbop["parts"]:
        f = parse_fastener(p["name"])
        if f:
            tool, tool_basis = tool_for(f)
            tq = torque_for(f, mat_names, locknut_threads, resin_seed)
            fam = FAMILY[f["family"]]
            fasteners.append(dict(
                name=p["name"], qty=p["qty"], family=f["family"], thread=f["thread"] or "待确认",
                length=f["length"] or "—", head=fam["head"], standard=fam["std"], standard_source=fam["std_src"],
                drive=fam["drive"], tool=tool, tool_basis=tool_basis,
                torque_class=tq["cls"], torque_nm=tq["value"], torque_text=tq["text"], torque_why=tq["why"],
            ))
            FAST_TOTAL[p["name"]] += p["qty"] or 0
            if tool not in ("无（装配辅件）",):
                add_tool(tool, tool_basis, op)
        else:
            parts.append(dict(name=p["name"], qty=p["qty"]))
    FAST_BY_OP[op] = fasteners

    # 官方工具栏
    for t in wbop.get("tools") or []:
        add_tool(t["name"] + "（官方）", "官方 Workbook 工具栏", op)

    # --- 功能序 ---
    s7 = S7["ops"].get(op)
    if s7:
        seq_names = [OCC2NAME.get(i["occ"], i["occ"]) for i in s7["order"]]
        fs = s7.get("functional_sort") or {}
        seed = OCC2NAME.get(fs.get("seed"), fs.get("seed"))
        unsupported = [OCC2NAME.get(x, x) for x in lst(fs.get("unsupported"))]
        seq = [dict(name=n, qty=q) for n, q in runs(seq_names)]
        seq_note = (f"来源：S7-工序内播放序 v3 functional_sort（{fs.get('rule','')}）；"
                    f"工序内硬约束 {fs.get('intra_constraints_applied', 0)} 条")
    else:
        seq, seed, unsupported = [], None, []
        seq_note = "本步在 S7 v3 中无实体动件（接线/管路/或零件全部并入相邻步的网格），承载先后按 Workbook 零件行序与功能线判读"

    # --- 功能线内容 ---
    subs = SUB_BY_OP.get(op, [])
    fn_names = [s.get("name") for s in subs if s.get("name")]
    install_dirs, adjustments, fn_notes, fn_funcs, fn_prins, fn_deps = [], [], [], [], [], []
    fn_howtos: list[str] = []
    for s in subs:
        im = s.get("install_method") or {}
        if im.get("direction"):
            install_dirs.append(f"{s['_track']}｜{im['direction']}")
        if im.get("adjustment_after_mount"):
            adjustments.append(dict(text=im["adjustment_after_mount"], src=f"功能线 {s['_track']}·{s.get('name','')}"))
        if im.get("notes"):
            fn_notes.append(f"{s['_track']}｜{im['notes']}")
        if s.get("function"):
            fn_funcs.append(f"{s['_track']}｜{s['function']}")
        if s.get("principle"):
            fn_prins.append(f"{s['_track']}｜{s['principle']}")
        for dep in lst(s.get("functional_dependencies")):
            fn_deps.append(f"{s['_track']}｜{dep}")
        for h in lst(im.get("how_tos")):
            if h not in fn_howtos:
                fn_howtos.append(h)
    for h in lst(wbop.get("how_tos")):
        if h not in fn_howtos:
            fn_howtos.append(h)

    for c in ADJ_BY_OP.get(op, []):
        adjustments.append(dict(text=f"本步之后、{c['to']} 之前：{c['why']}",
                                src=f"顺序约束 ADJUST_AFTER（{c['track']} 线，{c['source']}，置信 {c['confidence']}）"))
    for c in CONSTRAINTS:
        if c["kind"] == "ADJUST_AFTER" and not c["intra"] and c["to"] == op:
            adjustments.append(dict(text=f"开始本步前须确认 {c['frm']} 之后的调整已完成：{c['why']}",
                                    src=f"顺序约束 ADJUST_AFTER（{c['track']} 线）"))

    # 设计规则（来自功能线单元）
    design_rules = []
    seen_rule = set()
    for u in UNIT_BY_OP.get(op, []):
        for r in lst(u.get("design_rules")):
            key = r.get("rule")
            if key and key not in seen_rule:
                seen_rule.add(key)
                design_rules.append(dict(rule=key, basis=r.get("basis"), confidence=r.get("confidence"),
                                         unit=u["unit_id"], track=u["_track"]))

    # --- 前置条件 ---
    pres = []
    seen_pre = set()
    for c in PRE_BY_OP.get(op, []):
        k = (c["frm"], c["kind"], c["why"][:40])
        if k in seen_pre:
            continue
        seen_pre.add(k)
        pres.append(dict(op=c["frm"], title=next((w["title"] for w in WB["operations"] if w["op"] == c["frm"]), ""),
                         kind=c["kind"], why=c["why"], source=c["source"], confidence=c["confidence"], track=c["track"]))
    pres.sort(key=lambda x: (OP_INDEX.get(x["op"], 999), x["kind"]))
    doc_pred = [dict(op=p, title=next((w["title"] for w in WB["operations"] if w["op"] == p), ""),
                     kind="手册序", why=f"S3 工序 DAG（{s3.get('evidence_class','')}）", source="官方资料",
                     confidence="高", track="S3")
                for p in lst(s3.get("predecessors"))]

    hints = []
    seen_h = set()
    for c in HINT_BY_OP.get(op, []):
        k = (c["kind"], c["frm"], c["to"], c["why"][:40])
        if k in seen_h:
            continue
        seen_h.add(k)
        hints.append(dict(kind=c["kind"], frm=c["frm"], to=c["to"], why=c["why"],
                          source=c["source"], confidence=c["confidence"], track=c["track"]))

    intra = [dict(why=c["why"], from_part=c["from_part"], to_part=c["to_part"], parts=c["parts"],
                  kind=c["kind"], source=c["source"], confidence=c["confidence"], track=c["track"])
             for c in INTRA_BY_OP.get(op, [])]

    # --- 待确认问题 ---
    questions = []
    for q in lst(wbop.get("problems")):
        questions.append(dict(text=q, src="官方 Workbook Problems & questions 栏（原文）"))
    for u in UNIT_ANCHOR.get(op, []):
        scope = "、".join(ops_in_code("、".join(str(x) for x in lst(u.get("workbook_ops"))))[:6]) or op
        for q in lst(u.get("open_questions")):
            questions.append(dict(text=q, src=f"功能线 {u['_track']}·{u['unit_id']}（覆盖 {scope} 等）open_questions"))
        for g in lst(u.get("data_gaps")):
            questions.append(dict(text="资料缺口：" + str(g), src=f"功能线 {u['_track']}·{u['unit_id']} data_gaps"))
    fam_used = {f["family"] for f in fasteners}
    if "HF-screw" in fam_used:
        questions.append(dict(text="HF-screw 无标准号与头型记录（全机 284 根，用量第一），本规程按 DIN 6921 六角法兰头推断工具与扭矩，需厂家确认。",
                              src="本规程推断"))
    if "Thermoplastic Screw" in fam_used:
        questions.append(dict(text="Thermoplastic Screw 头型无记录，本规程假定内六角头，需确认（若为梅花槽则工具不同）。", src="本规程推断"))
    if "Standoff" in fam_used:
        questions.append(dict(text="Standoff FF 8mm 只给长度未给螺纹规格，工具与扭矩待确认。", src="本规程推断"))
    if "Set Screw" in fam_used:
        questions.append(dict(text="Set Screw 未给螺纹规格，本规程按 M5 顶丝取值，需确认。", src="本规程推断"))
    if op == "WB-01.1":
        questions.append(dict(
            text="官方工具栏 Allen Key 6 与 Wrench 14 同 ISO 7380／DIN 985 标准对边表不一致（M8 半圆头应为 5 mm 内六角、M8 尼龙锁母应为 13 mm 对边），两处均大一档；需确认 B-screw 实际头型与 Lock Nut 实际对边。",
            src="本规程核对"))

    # 去重
    _seen = set()
    questions = [q for q in questions if not (q["text"] in _seen or _seen.add(q["text"]))]
    _seen = set()
    adjustments = [a for a in adjustments
                   if not re.match(r"^\s*无(（|\(|$|。|；)", a["text"])
                   and not (a["text"] in _seen or _seen.add(a["text"]))]

    # 计数对账
    mism = [r for r in lst(s3.get("part_rows")) if r.get("status") != "COUNT_MATCH"]

    resin_note = None
    if resin_parts_here and not resin_seed:
        resin_note = ("本步含树脂 3D 打印件（" + "、".join(resin_parts_here[:6])
                      + "）；直接压在树脂件上的螺钉应降至 C 类值（M3 0.4 / M4 0.9 / M5 1.8 / M6 3.0 N·m，推断），"
                        "并建议加大垫圈分散压强。")

    CARDS.append(dict(
        op=op, step=wbop["step"], chapter=wbop["chapter"], title=wbop["title"],
        title_zh="；".join(fn_names) if fn_names else None,
        bench_group=bench, bench_role=role, join_op=join_op, joins_here=joins_here,
        is_wiring=op in WIRING_OPS,
        glb_root_node=wbop.get("glb_root_node"),
        functional_order=seq, seed_part=seed, unsupported_parts=unsupported, functional_order_note=seq_note,
        parts=parts, fasteners=fasteners, resin_seed=resin_seed,
        resin_parts=resin_parts_here, resin_note=resin_note,
        official_tools=[t["name"] for t in (wbop.get("tools") or [])],
        how_tos=fn_howtos,
        function=fn_funcs, principle=fn_prins, install_direction=install_dirs,
        functional_dependencies=fn_deps,
        adjustments=adjustments, design_rules=design_rules,
        prerequisites=pres, doc_predecessors=doc_pred, hints=hints, intra_constraints=intra,
        official_notes=wbop.get("notes"), official_remarks=lst(wbop.get("remarks")),
        official_problems=lst(wbop.get("problems")),
        function_notes=fn_notes,
        open_questions=questions,
        count_mismatch=[dict(name=r["sop_name"], sop_qty=r["sop_qty"], glb=r.get("glb_meshes"),
                             bound=r.get("bound_occurrences"), status=r.get("status")) for r in mism],
    ))

# =====================================================================================
# 4. 汇总表
# =====================================================================================

FASTENER_SUMMARY = []
for name, qty in sorted(FAST_TOTAL.items(), key=lambda kv: (-kv[1], kv[0])):
    f = parse_fastener(name)
    fam = FAMILY[f["family"]]
    tool, tool_basis = tool_for(f)
    ops_using = [c["op"] for c in CARDS if any(x["name"] == name for x in c["fasteners"])]
    tqs = sorted({x["torque_text"] for c in CARDS for x in c["fasteners"] if x["name"] == name})
    clss = sorted({x["torque_class"] for c in CARDS for x in c["fasteners"] if x["name"] == name})
    FASTENER_SUMMARY.append(dict(
        name=name, qty=qty, family=f["family"], thread=f["thread"] or "待确认", length=f["length"] or "—",
        head=fam["head"], standard=fam["std"], standard_source=fam["std_src"], drive=fam["drive"],
        tool=tool, tool_basis=tool_basis, torque_classes=clss, torque_values=tqs,
        n_ops=len(ops_using), ops=ops_using,
    ))

FAMILY_SUMMARY = []
_famq = Counter()
for name, qty in FAST_TOTAL.items():
    _famq[parse_fastener(name)["family"]] += qty
for fam, qty in sorted(_famq.items(), key=lambda kv: -kv[1]):
    FAMILY_SUMMARY.append(dict(family=fam, qty=qty, head=FAMILY[fam]["head"], standard=FAMILY[fam]["std"],
                               drive=FAMILY[fam]["drive"], source=FAMILY[fam]["std_src"]))

TOOL_SUMMARY = sorted(TOOL_TOTAL.values(), key=lambda r: (-len(r["ops"]), r["name"]))
for r in TOOL_SUMMARY:
    r["n_ops"] = len(r["ops"])

TORQUE_TABLE = []
for th in ("M3", "M4", "M5", "M6", "M8", "M10", "M12"):
    a = TORQUE_A[th]
    TORQUE_TABLE.append(dict(
        thread=th, A=rnd(a), B=rnd(a * K_ALU), C=rnd(a * K_RESIN),
        F=rnd(a + PREVAIL[th]), prevail=f"{PREVAIL[th]:.1f}",
    ))

# ---- 台架并行计划 ----
BENCH_PLAN = []
for gid, gops in BENCH_OPS.items():
    join = JOIN_OF_BENCH.get(gid)
    join_idx = OP_INDEX.get(join, 999) if join else 999
    ext_pre, late_pre = [], []
    for gop in gops:
        for c in PRE_BY_OP.get(gop, []):
            if c["frm"] in gops:
                continue
            (late_pre if OP_INDEX.get(c["frm"], 999) >= join_idx else ext_pre).append(c["frm"])
    ext_pre = sorted(set(ext_pre), key=lambda x: OP_INDEX.get(x, 999))
    late_pre = sorted(set(late_pre), key=lambda x: OP_INDEX.get(x, 999))
    earliest = max([OP_INDEX.get(x, -1) for x in ext_pre], default=-1)
    n_parts = sum(len(c["parts"]) + len(c["fasteners"]) for c in CARDS if c["op"] in gops)
    n_qty = sum((p["qty"] or 0) for c in CARDS if c["op"] in gops for p in c["parts"] + c["fasteners"])
    BENCH_PLAN.append(dict(
        group=gid, ops=gops, n_ops=len(gops), join_op=join,
        join_index=OP_INDEX.get(join, 999) if join else None,
        external_prereqs=ext_pre, late_prereqs=late_pre,
        earliest_start_after=(OPS_ORDER[earliest] if earliest >= 0 else "开工即可"),
        earliest_index=earliest,
        window=f"{OPS_ORDER[earliest] if earliest >= 0 else '开工'} → {join}",
        part_rows=n_parts, part_qty=n_qty,
        title="；".join(next(c["title"] for c in CARDS if c["op"] == g) for g in gops),
    ))
BENCH_PLAN.sort(key=lambda r: (r["earliest_index"], r["join_index"] if r["join_index"] is not None else 999))

# 并行波次：可同时开工的组（外部前置相同 / 均已满足）
WAVES: list[dict] = []
_remaining = list(BENCH_PLAN)
while _remaining:
    gate = _remaining[0]["earliest_index"]
    wave = [g for g in _remaining if g["earliest_index"] <= gate]
    latest_join = min(g["join_index"] for g in wave if g["join_index"] is not None)
    WAVES.append(dict(
        wave=len(WAVES) + 1,
        gate=(OPS_ORDER[gate] if gate >= 0 else "开工即可"),
        groups=[g["group"] for g in wave],
        must_finish_by=OPS_ORDER[latest_join],
        n_groups=len(wave),
        ops=[o for g in wave for o in g["ops"]],
    ))
    _remaining = [g for g in _remaining if g not in wave]

# ---- 约束总表 ----
CONSTRAINT_TABLE = [dict(track=c["track"], kind=c["kind"], frm=c["frm"], to=c["to"], scope=("工序内" if c["intra"] else "工序间"),
                         why=c["why"], source=c["source"], confidence=c["confidence"], status=c["status"])
                    for c in CONSTRAINTS]
PRE_COUNT = sum(len(c["prerequisites"]) for c in CARDS)

# ---- 待确认问题总表 ----
QUESTIONS = []
for c in CARDS:
    for q in c["open_questions"]:
        QUESTIONS.append(dict(op=c["op"], title=c["title"], text=q["text"], src=q["src"]))

# ---- 装后调整总表 ----
ADJUSTMENTS = []
for c in CARDS:
    for a in c["adjustments"]:
        ADJUSTMENTS.append(dict(op=c["op"], title=c["title"], text=a["text"], src=a["src"]))

# =====================================================================================
# 5. sop.json
# =====================================================================================

DOC = OrderedDict(
    [
        ("$schema", "assembly-ontology.process-spec/v1"),
        ("record_id", "OLSK-V3-PROCESS-SPEC-V001"),
        ("title", "OLSK Large CNC V3 装配工艺规程"),
        ("version", VERSION),
        ("generated_at", datetime.datetime.now().astimezone().isoformat()),
        ("purpose", "把官方 Workbook 的 114 步升级为可执行工艺卡：功能序、零件与紧固件清单、工具与推荐扭矩、装后调整与检验点、前置条件、官方原文与待确认问题。"),
        ("machine", dict(
            name="OLSK Large CNC V3",
            milling_area="2500 mm × 1250 mm",
            travel=dict(X_mm=float(SETTINGS.get("$130", 0)), Y_mm=float(SETTINGS.get("$131", 0)), Z_mm=float(SETTINGS.get("$132", 0))),
            steps_per_mm=dict(X=float(SETTINGS.get("$100", 0)), Y=float(SETTINGS.get("$101", 0)), Z=float(SETTINGS.get("$102", 0))),
            max_rate_mm_min=dict(X=float(SETTINGS.get("$110", 0)), Y=float(SETTINGS.get("$111", 0)), Z=float(SETTINGS.get("$112", 0))),
            spindle_rpm_max=float(SETTINGS.get("$30", 0)),
            designer="InMachines Ingrassia GmbH",
        )),
        ("sources", [dict(path=p, sha256=sha256(CASE / p)) for p in SRC_FILES]),
        ("claim_boundary", [
            "官方 Workbook 全表没有扭矩、螺纹胶、润滑、验收字段（field_semantics: torque/adhesive/acceptance 均为 UNKNOWN），114 步中只有 WB-01.1 填了工具栏。",
            "本规程的扭矩全部是按 ISO 898-1 8.8 级查表值与材料折减系数推算的参考值，不是官方值，投产前必须由设计方确认。",
            "工具规格由紧固件头型标准对边表推断；HF-screw、Thermoplastic Screw、Standoff、Set Screw 四类无标准号或规格记录，标为待确认。",
            "工序内先后取自 本体/S7-工序内播放序.v3.json 的 functional_sort（几何认证 + 承载优先 + 工序内硬约束），零件名用手册命名。",
            "前置条件取自 功能属性/tracks/*/order_constraints.json 的 MUST_BEFORE/ACCESS/WIRE_BEFORE_CLOSE，逐条保留原始来源与置信度；SAFETY/SAME_BENCH/MUST_AFTER 另列为提示，其中 MUST_AFTER 各线书写方向不一致，未纳入前置条件推导。",
            "官方 Notes/Remarks/Problems 逐字保留英文原文，不做归一与改写。",
        ]),
        ("summary", dict(
            operations=len(CARDS), chapters=len({c["chapter"] for c in CARDS}),
            bench_groups=len(BENCH_OPS), bench_ops=sum(len(v) for v in BENCH_OPS.values()),
            join_ops=len(JOIN_OF_BENCH), wiring_ops=len(WIRING_OPS),
            part_rows=sum(len(c["parts"]) for c in CARDS),
            part_qty=sum((p["qty"] or 0) for c in CARDS for p in c["parts"]),
            fastener_rows=sum(len(c["fasteners"]) for c in CARDS),
            fastener_qty=sum(FAST_TOTAL.values()),
            fastener_specs=len(FAST_TOTAL),
            tools=len(TOOL_SUMMARY),
            prerequisites=PRE_COUNT,
            constraints_total=len(CONSTRAINTS),
            adjustments=len(ADJUSTMENTS),
            open_questions=len(QUESTIONS),
            ops_with_functional_order=sum(1 for c in CARDS if c["functional_order"]),
        )),
        ("torque_reference", dict(
            basis="ISO 898-1 8.8 级钢制螺纹副，总摩擦系数 μ = 0.14，取 90% 屈服利用率的常用查表值（N·m）",
            disclaimer="参考值，非官方值。官方 Workbook 无任何扭矩字段。",
            classes=TORQUE_CLASS, base_values=TORQUE_A,
            k_aluminium=K_ALU, k_resin=K_RESIN,
            prevailing_torque=PREVAIL,
            cross_check="铝型材厂家（item / Bosch Rexroth 等）对 M5/M6/M8 型材连接常见推荐 5 / 10 / 20~25 N·m，与本表 B 类折减值 4.1 / 7.1 / 17 N·m 同量级，本表偏保守。",
            table=TORQUE_TABLE,
        )),
        ("fastener_families", FAMILY_SUMMARY),
        ("fastener_totals", FASTENER_SUMMARY),
        ("tool_totals", TOOL_SUMMARY),
        ("bench_parallel_plan", dict(groups=BENCH_PLAN, waves=WAVES)),
        ("how_tos", list(HOWTO.values())),
        ("constraints", CONSTRAINT_TABLE),
        ("adjustments", ADJUSTMENTS),
        ("open_questions", QUESTIONS),
        ("operations", CARDS),
    ]
)
OUT_JSON.write_text(json.dumps(DOC, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"[1/3] sop.json  {OUT_JSON.stat().st_size/1024:.0f} KB  工序 {len(CARDS)}")

# =====================================================================================
# 6. xlsx
# =====================================================================================

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(size=9)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, Alignment(wrap_text=True, vertical="center"), BORDER
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = WRAP if (isinstance(c.value, str) and len(c.value) > 18) else TOP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"
    ws.row_dimensions[1].height = 30
    return ws


wb = Workbook()
wb.remove(wb.active)

# 0 封面/说明
ws = wb.create_sheet("00 说明")
info = [
    ("文件", f"OLSK Large CNC V3 装配工艺规程 {VERSION}"),
    ("生成日期", DATE),
    ("工序数", len(CARDS)),
    ("零件行数（非紧固件）", sum(len(c["parts"]) for c in CARDS)),
    ("紧固件总数", sum(FAST_TOTAL.values())),
    ("紧固件规格数", len(FAST_TOTAL)),
    ("台架预装组", f"{len(BENCH_OPS)} 组 / {sum(len(v) for v in BENCH_OPS.values())} 步"),
    ("工序间前置条件", PRE_COUNT),
    ("顺序约束总条数", len(CONSTRAINTS)),
    ("装后调整/检验点", len(ADJUSTMENTS)),
    ("待确认问题", len(QUESTIONS)),
    ("扭矩依据", DOC["torque_reference"]["basis"]),
    ("扭矩声明", DOC["torque_reference"]["disclaimer"]),
]
for k, v in info:
    ws.append([k, v])
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 110
for row in ws.iter_rows():
    row[0].font = Font(bold=True, size=10)
    row[1].font = BODY_FONT
    row[1].alignment = WRAP
ws.append([])
ws.append(["声明", "本规程的扭矩、工具规格、材料折减均为推断参考值，不是官方数据；官方 Workbook 无扭矩/胶粘/验收字段。"])
ws.append(["输入", "；".join(SRC_FILES[:8]) + " 等 " + str(len(SRC_FILES)) + " 个只读文件（sha256 见 sop.json）"])

# 1 工艺卡总表
rows = []
for c in CARDS:
    rows.append([
        c["op"], c["step"], c["chapter"], c["title"], c["title_zh"] or "", c["bench_role"],
        fmt_runs([(x["name"], x["qty"]) for x in c["functional_order"]]) or "（无实体动件）",
        c["seed_part"] or "",
        "；".join(f"{p['name']}×{p['qty']}" for p in c["parts"]),
        "；".join(f"{f['name']}×{f['qty']}" for f in c["fasteners"]),
        sum((f["qty"] or 0) for f in c["fasteners"]),
        "；".join(sorted({f["tool"] for f in c["fasteners"]} | set(c["official_tools"]))),
        "；".join(f"{f['name']}: {f['torque_text']}" for f in c["fasteners"] if f["torque_text"] not in ("—",)),
        c.get("resin_note") or "",
        "；".join(a["text"] for a in c["adjustments"]),
        "；".join(f"{p['op']}({p['kind']})" for p in c["prerequisites"]),
        c["official_notes"] or "",
        " / ".join(c["official_remarks"]),
        " / ".join(c["official_problems"]),
        "、".join(c["how_tos"]),
        len(c["open_questions"]),
    ])
sheet(wb, "01 工艺卡总表",
      ["工序", "步号", "章", "标题(原文)", "中文部套名", "台架/上机角色", "功能序（承载先后）", "承载基准件",
       "零件清单", "紧固件清单", "紧固件数", "工具", "推荐扭矩(参考值)", "树脂件提示", "装后调整/检验点", "前置条件",
       "官方 Notes", "官方 Remarks", "官方 Problems", "How-To", "待确认"],
      rows, [10, 8, 6, 30, 24, 26, 46, 20, 40, 36, 9, 30, 34, 40, 44, 22, 40, 40, 26, 10, 8])

# 2 零件清单
rows = [[c["op"], c["step"], c["title"], i + 1, p["name"], p["qty"]]
        for c in CARDS for i, p in enumerate(c["parts"])]
sheet(wb, "02 零件清单", ["工序", "步号", "标题", "序", "零件名(原文)", "数量"], rows, [10, 8, 30, 6, 46, 8])

# 3 紧固件明细
rows = [[c["op"], c["step"], f["name"], f["qty"], f["family"], f["thread"], f["length"], f["head"],
         f["standard"], f["drive"], f["tool"], f["tool_basis"], f["torque_class"], f["torque_text"], f["torque_why"]]
        for c in CARDS for f in c["fasteners"]]
sheet(wb, "03 紧固件明细",
      ["工序", "步号", "规格(原文)", "数量", "家族", "螺纹", "长度", "头型", "标准号(推断)", "驱动",
       "工具(推断)", "工具依据", "扭矩类别", "推荐扭矩(参考值)", "扭矩判定理由"],
      rows, [10, 8, 26, 8, 18, 8, 9, 24, 34, 10, 30, 30, 9, 20, 46])

# 4 紧固件总表
rows = [[r["name"], r["qty"], r["family"], r["thread"], r["length"], r["head"], r["standard"], r["standard_source"],
         r["drive"], r["tool"], "/".join(r["torque_classes"]), "；".join(r["torque_values"]), r["n_ops"],
         "、".join(r["ops"][:14]) + (" …" if len(r["ops"]) > 14 else "")]
        for r in FASTENER_SUMMARY]
sheet(wb, "04 紧固件总表",
      ["规格(原文)", "总数", "家族", "螺纹", "长度", "头型", "标准号(推断)", "标准来源", "驱动", "工具(推断)",
       "扭矩类别", "推荐扭矩(参考值)", "涉及工序数", "涉及工序"],
      rows, [26, 8, 18, 8, 9, 24, 34, 10, 10, 30, 10, 24, 10, 50])

# 4b 家族汇总
rows = [[r["family"], r["qty"], r["head"], r["standard"], r["drive"], r["source"]] for r in FAMILY_SUMMARY]
rows.append(["合计", sum(FAST_TOTAL.values()), "", "", "", ""])
sheet(wb, "05 紧固件家族", ["家族", "总数", "头型", "标准号", "驱动", "标准来源"], rows, [22, 10, 26, 46, 12, 12])

# 6 工具总表
rows = [[r["name"], r["n_ops"], r["basis"], "、".join(r["ops"][:20]) + (" …" if len(r["ops"]) > 20 else "")]
        for r in TOOL_SUMMARY]
sheet(wb, "06 工具总表", ["工具", "涉及工序数", "依据", "涉及工序"], rows, [40, 10, 44, 70])

# 7 扭矩参考表
rows = [[t["thread"], t["A"], t["B"], t["C"], t["prevail"], t["F"]] for t in TORQUE_TABLE]
ws = sheet(wb, "07 扭矩参考表",
           ["螺纹", "A 钢制螺纹副 (N·m)", f"B 铝件 ×{K_ALU:.2f} (N·m)", f"C 树脂件 ×{K_RESIN:.2f} (N·m)",
            "锁母自锁力矩 (N·m)", "F 锁紧螺母合计 (N·m)"],
           rows, [10, 20, 20, 20, 18, 22])
ws.append([])
ws.append(["依据", DOC["torque_reference"]["basis"]])
ws.append(["声明", DOC["torque_reference"]["disclaimer"]])
ws.append(["交叉核对", DOC["torque_reference"]["cross_check"]])
for k, v in TORQUE_CLASS.items():
    ws.append([k, v["name"] + "｜" + v["rule"]])
for r in ws.iter_rows(min_row=len(rows) + 2):
    r[0].font = Font(bold=True, size=9)
    if len(r) > 1:
        r[1].font = BODY_FONT
        r[1].alignment = WRAP

# 8 前置约束
rows = [[c["op"], c["title"], p["op"], p["title"], p["kind"], p["why"], p["source"], p["confidence"], p["track"]]
        for c in CARDS for p in c["prerequisites"]]
sheet(wb, "08 前置条件", ["工序", "标题", "前置工序", "前置标题", "约束类型", "理由", "来源", "置信", "功能线"],
      rows, [10, 28, 10, 28, 18, 76, 12, 8, 8])

# 9 全部顺序约束
rows = [[c["track"], c["kind"], c["frm"], c["to"], c["scope"], c["why"], c["source"], c["confidence"], c["status"] or ""]
        for c in CONSTRAINT_TABLE]
sheet(wb, "09 顺序约束总表", ["功能线", "类型", "from", "to", "范围", "理由", "来源", "置信", "影片状态"],
      rows, [8, 20, 10, 10, 10, 84, 12, 8, 16])

# 10 装后调整/检验点
rows = [[a["op"], a["title"], a["text"], a["src"]] for a in ADJUSTMENTS]
sheet(wb, "10 装后调整与检验", ["工序", "标题", "调整/检验内容", "来源"], rows, [10, 28, 90, 40])

# 11 台架并行计划
rows = [[g["group"], g["n_ops"], "、".join(g["ops"]), g["join_op"], g["earliest_start_after"],
         "、".join(g["external_prereqs"]) or "无", g["part_rows"], g["part_qty"], g["title"]]
        for g in BENCH_PLAN]
sheet(wb, "11 台架并行计划",
      ["台架组", "步数", "包含工序", "上机步", "最早可开工（此步后）", "外部前置", "行数", "件数", "工序标题"],
      rows, [10, 6, 34, 10, 20, 30, 8, 8, 70])

rows = [[w["wave"], w["gate"], w["n_groups"], "、".join(w["groups"]), w["must_finish_by"], "、".join(w["ops"])]
        for w in WAVES]
sheet(wb, "12 台架并行波次", ["波次", "开工条件（此步完成后）", "组数", "台架组", "最早需交付于", "包含工序"],
      rows, [8, 22, 8, 40, 16, 70])

# 13 功能序
rows = [[c["op"], c["step"], c["title"], fmt_runs([(x["name"], x["qty"]) for x in c["functional_order"]]) or "（无实体动件）",
         c["seed_part"] or "", "、".join(dict.fromkeys(c["unsupported_parts"])), c["functional_order_note"]]
        for c in CARDS]
sheet(wb, "13 功能序", ["工序", "步号", "标题", "工序内安装先后（手册零件名）", "承载基准件", "无支撑件（提示）", "说明"],
      rows, [10, 8, 28, 88, 22, 34, 46])

# 14 待确认问题
rows = [[q["op"], q["title"], q["text"], q["src"]] for q in QUESTIONS]
sheet(wb, "14 待确认问题", ["工序", "标题", "问题", "来源"], rows, [10, 28, 92, 40])

# 15 How-To
rows = [[h["id"], h["title"], "；".join(h["parts"]), "；".join(h["tools"]), " ".join(h["remarks"])]
        for h in HOWTO.values()]
sheet(wb, "15 How-To 原文", ["编号", "标题", "零件", "工具", "要点原文"], rows, [10, 34, 34, 34, 100])

wb.save(OUT_XLSX)
print(f"[2/3] xlsx      {OUT_XLSX.stat().st_size/1024:.0f} KB  {len(wb.sheetnames)} 个工作表")

# =====================================================================================
# 7. LaTeX → PDF
# =====================================================================================

REP = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
    "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
    "→": r"$\rightarrow$", "←": r"$\leftarrow$", "↔": r"$\leftrightarrow$", "⇒": r"$\Rightarrow$",
    "≈": r"$\approx$", "≤": r"$\le$", "≥": r"$\ge$", "≠": r"$\ne$", "×": r"$\times$",
    "±": r"$\pm$", "°": r"$^\circ$", "℃": "°C", "µ": r"$\mu$", "μ": r"$\mu$", "Δ": r"$\Delta$",
    "•": r"\textbullet{}", "√": r"$\surd$", "≡": r"$\equiv$", "∅": "Ø", "²": r"$^{2}$", "³": r"$^{3}$",
    "′": r"$^{\prime}$", "″": r"$^{\prime\prime}$", "✅": "[√]", "❌": "[×]", "⚠": "[!]",
    "①": "(1)", "②": "(2)", "③": "(3)", "④": "(4)", "⑤": "(5)", "⑥": "(6)", "⑦": "(7)", "⑧": "(8)",
    "≥": r"$\ge$", "…": r"\ldots{}",
}


def tex(s) -> str:
    if s is None:
        return ""
    return "".join(REP.get(ch, ch) for ch in str(s))


ID_RE = re.compile(r"(WB-\d+(?:\.\d+)?|BENCH-\d+|M\d{1,2}-\d{1,3})")


def texb(s) -> str:
    """tex() + 长零件名处允许断行；工序号/台架号/螺钉规格保持不断行。"""
    out = []
    for part in ID_RE.split(tex(s)):
        if not part:
            continue
        if ID_RE.fullmatch(part):
            out.append("\\mbox{" + part + "}")
        else:
            out.append(re.sub(r"(\\_|-|\.|/|（|\(|×|,)", r"\1\\allowbreak{}", part))
    return "".join(out)


def fit_spec(spec: str, total_mm: float = 158.0) -> str:
    pw = [float(x) for x in re.findall(r"p\{([0-9.]+)mm\}", spec)]
    rest = re.sub(r">\{[^}]*\}p\{[0-9.]+mm\}|p\{[0-9.]+mm\}", "", spec)
    n_lrc = len(re.findall(r"[lrc]", rest))
    ncol = len(pw) + n_lrc
    avail = total_mm - 2.2 * ncol - 13.0 * n_lrc
    s = sum(pw)
    if s > 0 and s > avail:
        f = avail / s
        spec = re.sub(r"p\{([0-9.]+)mm\}", lambda m: f"p{{{float(m.group(1)) * f:.1f}mm}}", spec)
    return spec


def longtable(spec, header, rows, caption=None):
    spec = fit_spec(spec)
    out = ["\\begingroup\\zihao{6}\\begin{longtable}{" + spec + "}"]
    if caption:
        out.append("\\caption{" + tex(caption) + "}\\\\")
    hd = " & ".join("\\textbf{" + tex(h) + "}" for h in header) + "\\\\ \\midrule"
    out.append("\\toprule " + hd + " \\endfirsthead")
    out.append("\\toprule " + hd + " \\endhead")
    out.append("\\bottomrule \\endfoot")
    for r in rows:
        out.append(" & ".join(texb(c) if c is not None else "" for c in r) + "\\\\")
    out.append("\\end{longtable}\\endgroup")
    return "\n".join(out)


def itemize(items, small=True):
    if not items:
        return ""
    o = ["\\begingroup\\zihao{6}" if small else "", "\\begin{itemize}[leftmargin=1.4em,itemsep=1pt,topsep=2pt]"]
    for it in items:
        o.append("\\item " + texb(it))
    o.append("\\end{itemize}")
    o.append("\\endgroup" if small else "")
    return "\n".join(x for x in o if x)


PREAMBLE = r"""% !TeX program = xelatex
\documentclass[UTF8,zihao=-4]{ctexart}
\usepackage[a4paper,top=22mm,bottom=20mm,left=20mm,right=20mm,headsep=7mm,footskip=11mm]{geometry}
\usepackage{fontspec}
\setmainfont{texgyretermes-regular.otf}[BoldFont=texgyretermes-bold.otf,ItalicFont=texgyretermes-italic.otf,BoldItalicFont=texgyretermes-bolditalic.otf]
\setsansfont{texgyreheros-regular.otf}[BoldFont=texgyreheros-bold.otf]
\setmonofont{texgyrecursor-regular.otf}[Scale=0.85]
\setCJKmainfont{Songti SC}[FakeStretch=0.94,AutoFakeBold=2.5]
\setCJKsansfont{Heiti SC}[FakeStretch=0.94,AutoFakeBold=2.5]
\setCJKmonofont{Heiti SC}[FakeStretch=0.94]
\usepackage{array,longtable,tabularx,booktabs,multirow,makecell,multicol}
\usepackage{graphicx,float,caption,xcolor,enumitem,fancyhdr,titlesec,indentfirst,ragged2e}
\PassOptionsToPackage{hyphens}{url}\usepackage[hidelinks]{hyperref}\urlstyle{tt}
\setlength{\parindent}{2em}\setlength{\parskip}{2pt}\setlength{\headheight}{15pt}\setlength{\emergencystretch}{3em}\hyphenpenalty=50
\setlength{\tabcolsep}{3pt}\renewcommand{\arraystretch}{1.15}
\captionsetup{font=small,labelfont=bf,skip=3pt}
\pagestyle{fancy}\fancyhf{}
\fancyhead[L]{\zihao{-5}OLSK Large CNC V3 装配工艺规程}\fancyhead[R]{\zihao{-5}""" + VERSION + r"""}
\fancyfoot[C]{\zihao{-5}第 \thepage\ 页}
\renewcommand{\headrulewidth}{0.4pt}\renewcommand{\footrulewidth}{0pt}
\titleformat{\section}{\centering\heiti\bfseries\zihao{3}}{\thesection}{0.6em}{}
\titleformat{\subsection}{\heiti\bfseries\zihao{4}}{\thesubsection}{0.6em}{}
\titleformat{\subsubsection}{\heiti\bfseries\zihao{-4}}{\thesubsubsection}{0.5em}{}
\titleformat{\paragraph}[hang]{\heiti\zihao{5}}{}{0pt}{}
\titlespacing*{\section}{0pt}{1.2\baselineskip}{0.8\baselineskip}
\titlespacing*{\subsection}{0pt}{0.8\baselineskip}{0.4\baselineskip}
\titlespacing*{\subsubsection}{0pt}{0.6\baselineskip}{0.3\baselineskip}
\titlespacing*{\paragraph}{0pt}{0.45\baselineskip}{0.2\baselineskip}
\setcounter{secnumdepth}{3}\setcounter{tocdepth}{2}
\definecolor{cardbg}{RGB}{242,245,250}
\begin{document}
"""

COVER = r"""
\begin{titlepage}
\centering
\vspace*{26mm}
{\heiti\zihao{1} OLSK Large CNC V3\par}
\vspace{4mm}
{\heiti\zihao{2} 装配工艺规程\par}
\vspace{6mm}
{\songti\zihao{3} 114 道工序工艺卡 · 紧固件与工具总表 · 扭矩参考表 · 台架并行计划\par}
\vspace{16mm}
{\zihao{-3} 版本 """ + VERSION + r""" \quad """ + DATE + r"""\par}
\vspace{26mm}
\begin{tabular}{r@{\hspace{1em}}p{112mm}}
\zihao{-4}机器 & \zihao{-4}OLSK Large CNC V3，加工范围 2500\,mm $\times$ 1250\,mm，设计方 InMachines Ingrassia GmbH\\[3pt]
\zihao{-4}官方依据 & \zihao{-4}\texttt{Workbook.csv}（114 步逐字搬运）、\texttt{HowTo\_Blad1.csv}、手册源码 preparing 分组\\[3pt]
\zihao{-4}本案依据 & \zihao{-4}本体 S3 工序实例绑定、S7 工序内播放序 v3、功能属性 A–H 线（含 403 条顺序约束）\\[3pt]
\zihao{-4}编制日期 & \zihao{-4}""" + DATE + r"""\\[3pt]
\zihao{-4}文档性质 & \zihao{-4}工艺文件草案；扭矩与工具规格为推断参考值，非官方数据\\
\end{tabular}
\vfill
\begin{minipage}{0.86\textwidth}\zihao{5}
\textbf{重要声明}\quad 官方 Workbook 全表\textbf{没有}扭矩、螺纹胶、润滑与验收字段，114 步中只有 WB-01.1 填写了工具。
本规程内所有扭矩数值均为按 ISO 898-1 8.8 级查表值与材料折减系数推算的\textbf{参考值}，所有工具规格均由紧固件头型标准对边表\textbf{推断}，
每一处均已在正文标注"推断"或"参考值"。投产前必须由设计方确认。
\end{minipage}
\vspace{10mm}
\end{titlepage}
\tableofcontents
\clearpage
"""


# -------------------------------------------------------------------------------------
# 转义纪律：longtable 的表头/表体/caption 与 itemize 的 item 一律接收「原始文本」，
# 由 longtable / itemize 内部统一 tex()/texb()。函数外拼接的裸 LaTeX 命令不经过 tex()。
# -------------------------------------------------------------------------------------


def sec_intro():
    o = ["\\section{适用范围、编制依据与声明}"]
    o.append("\\subsection{适用范围}")
    o.append(tex("本规程覆盖 OLSK Large CNC V3 整机装配的全部 114 道工序（Workbook 步号 01.1 至 78），"
                 "包括机架与底座、Y 轴与肩部、X/Z 轴与主轴头、床身与刀库、电气与控制、气动系统、外壳门窗共七条装配线，"
                 f"以及 {len(BENCH_OPS)} 个台架预装组（共 {sum(len(v) for v in BENCH_OPS.values())} 步）"
                 f"与 {len(WIRING_OPS)} 个接线/管路步。"))
    m = DOC["machine"]
    o.append("\\subsection{机器概况}")
    o.append(longtable("p{34mm}p{118mm}", ["项目", "数值"], [
        ["加工范围", "2500 mm × 1250 mm（README 规格）"],
        ["行程（控制器 $130/$131/$132）",
         f"机器 X {m['travel']['X_mm']:.0f} mm、Y {m['travel']['Y_mm']:.0f} mm、Z {m['travel']['Z_mm']:.0f} mm"],
        ["快移上限（$110/$111/$112）",
         f"X {m['max_rate_mm_min']['X']:.0f}、Y {m['max_rate_mm_min']['Y']:.0f}、Z {m['max_rate_mm_min']['Z']:.0f} mm/min"],
        ["主轴", f"4.5 kW BT30 自动换刀电主轴，最高 {m['spindle_rpm_max']:.0f} r/min（$30）"],
        ["传动", "三轴滚珠丝杠；Y 轴为旋转螺母，双电机同步"],
        ["机架", "CNC 铣孔铝方管 100×100×5 / 100×50×5 mm，连接板 + HF-screw M12-20 端面对接，铝角件加强；全机不使用 T 型螺母"],
        ["外壳", "全封闭，三面可开启（气弹簧提升窗）"],
        ["刀库", "14 位；自动对刀、自动清刀、气动打刀"],
        ["设计方", "InMachines Ingrassia GmbH（开源，OLSK）"],
    ], "机器概况（据 README 与控制器设置）"))
    o.append("\\subsection{编制依据}")
    o.append(longtable("p{62mm}p{90mm}", ["输入文件", "用途"], [
        ["装配工艺/sop-olsk-workbook.v1.json", "官方 Workbook 114 步逐字搬运：步号、标题、零件与数量、工具、Notes/Remarks/Problems、How-To"],
        ["功能属性/inputs/Workbook.csv", "官方 SOP 原表（上表的源文件，用于校验）"],
        ["功能属性/inputs/HowTo_Blad1.csv", "How-To 原文（H1–H14）"],
        ["装配工艺/manual-bench-groups.v1.json", "手册源码 Model.jsx 内建的 preparing 台架分组与 wiring 步骤列表"],
        ["本体/S3-工序实例绑定.v1.json", "台架组 bench_groups、上机结合 unit_joins、手册序 predecessors、逐工序动件与计数对账"],
        ["本体/S7-工序内播放序.v3.json", "工序内功能序（承载优先 + 工序内硬约束的实体序）"],
        ["功能属性/tracks/A..H/function.json", "承载件、安装方式与方向、装后调整、功能依赖、设计规则、待确认问题"],
        ["功能属性/tracks/A..H/order_constraints.json", f"工序间与工序内顺序约束共 {len(CONSTRAINTS)} 条（含理由、来源、置信度）"],
        ["功能属性/inputs/OLSK_README.md、Settings.txt", "机器概况与控制器行程/速度设置"],
    ], "编制依据（全部为只读输入，sha256 记录于 sop.json）"))
    o.append("\\subsection{声明与使用限制}")
    o.append(itemize(DOC["claim_boundary"], small=False))
    o.append("\\subsection{工艺卡的读法}")
    o.append(itemize([
        "【功能序】给出本步内零件的安装先后，来自 S7 v3 的 functional_sort（几何认证序 + 承载件优先 + 工序内硬约束），"
        "零件名一律用手册英文原名，同名连续件折叠为“名称×数量”。",
        "【承载基准件】是本步的落座基准，其余件都直接或间接落在它上面。",
        "【紧固件】的头型、工具与扭矩全部由规格前缀推断：HF-screw = 六角法兰头 → 扳手；B-screw = 半圆头内六角；"
        "C-screw = 圆柱头内六角；Countersunk Screw = 沉头内六角；Lock Nut = 扳手。",
        "【推荐扭矩】是参考值，类别见附录 C：A 类钢制螺纹副，B 类夹铝件折减，C 类夹树脂件折减，D 类自攻入塑料，"
        "E 类顶丝，F 类尼龙锁紧螺母。",
        "【前置条件】只收 MUST_BEFORE / ACCESS / WIRE_BEFORE_CLOSE 三类工序间约束，逐条给出前置步号与理由；"
        "SAFETY / SAME_BENCH / MUST_AFTER 另列为“顺序提示”。",
        "【官方原文】栏内的英文一字未改，包括原文的拼写错误。",
    ], small=False))
    return "\n".join(o)


def sec_route():
    o = ["\\section{总体工艺路线}"]
    o.append("\\subsection{章节与工序清单}")
    rows = [[c["op"], c["title"], c["bench_role"], str(len(c["parts"])),
             str(sum((f["qty"] or 0) for f in c["fasteners"])), str(len(c["prerequisites"]))] for c in CARDS]
    o.append(longtable("p{13mm}p{62mm}p{44mm}rrr",
                       ["工序", "标题（原文）", "台架/上机角色", "零件行", "紧固件", "前置"], rows, "114 道工序总览"))
    o.append("\\subsection{台架预装组与上机结合点}")
    rows = [[g["group"], "、".join(g["ops"]), g["join_op"], g["earliest_start_after"],
             "、".join(g["external_prereqs"]) or "无", str(g["part_qty"])] for g in BENCH_PLAN]
    pct = sum(len(v) for v in BENCH_OPS.values()) / 114 * 100
    o.append(longtable("p{14mm}p{44mm}p{13mm}p{20mm}p{40mm}r",
                       ["台架组", "包含工序", "上机步", "最早可开工", "外部前置", "件数"], rows,
                       f"{len(BENCH_PLAN)} 个台架预装组（预装 {sum(len(v) for v in BENCH_OPS.values())} 步，占 114 步的 {pct:.0f}%）"))
    o.append("\\subsection{并行预装波次}")
    o.append(tex("按“外部前置全部满足即可开工、必须在上机步之前交付”划窗口，台架组可按下表分波并行。"
                 "同一波内各组互不依赖，可分配给不同工位同时进行。"))
    rows = [[str(w["wave"]), w["gate"], str(w["n_groups"]), "、".join(w["groups"]), w["must_finish_by"]] for w in WAVES]
    o.append(longtable("rp{24mm}rp{62mm}p{22mm}",
                       ["波次", "开工条件（此步完成后）", "组数", "台架组", "最早交付节点"], rows, "台架并行波次"))
    o.append("\\subsection{接线与管路步}")
    o.append(tex("手册源码把下列步骤标为 wiring："
                 + "、".join(sorted(WIRING_OPS, key=lambda x: OP_INDEX.get(x, 999)))
                 + "。这些步骤在功能上受 WIRE_BEFORE_CLOSE 类约束支配——所有过线口、格兰头、线槽、拖链"
                   "必须在封盖之前装好并穿线。"))
    return "\n".join(o)


def card_tex(c: dict) -> str:
    o = ["\\clearpage",
         f"\\subsection*{{{tex(c['op'])}\\quad {tex(c['title'])}}}",
         f"\\addcontentsline{{toc}}{{subsection}}{{{tex(c['op'])} {tex(c['title'])}}}"]
    head = [
        ["工序号 / 步号 / 章", f"{c['op']} / {c['step']} / 第 {c['chapter']} 章"],
        ["中文部套名", c["title_zh"] or "（功能线未单列）"],
        ["台架/上机角色", c["bench_role"]],
    ]
    if c["joins_here"]:
        head.append(["接收台架完成件",
                     "、".join(c["joins_here"]) + "（来自 " + "、".join("+".join(BENCH_OPS[b]) for b in c["joins_here"]) + "）"])
    head.append(["承载基准件", c["seed_part"] or "（本步无实体动件）"])
    o.append(longtable("p{30mm}p{122mm}", ["项目", "内容"], head))

    o.append("\\paragraph{一、本步承载件与安装先后（功能序）}")
    if c["functional_order"]:
        o.append("\\begingroup\\zihao{6}\\noindent "
                 + texb(fmt_runs([(x["name"], x["qty"]) for x in c["functional_order"]])) + "\\par\\endgroup")
        if c["unsupported_parts"]:
            o.append("\\begingroup\\zihao{6}\\noindent 提示："
                     + texb("、".join(dict.fromkeys(c["unsupported_parts"])))
                     + tex(" 在几何上未找到已在位的支承面（可能靠工装/人手扶持，或由紧固件夹持），装配时需临时支撑。")
                     + "\\par\\endgroup")
    else:
        o.append("\\begingroup\\zihao{6}\\noindent " + tex(c["functional_order_note"]) + "\\par\\endgroup")
    if c["install_direction"]:
        o.append(itemize(["安装方向：" + d for d in c["install_direction"]]))
    if c["intra_constraints"]:
        rows = [[x["kind"], x["from_part"] or "、".join(x["parts"][:2]), x["to_part"] or "", x["why"],
                 f"{x['source']}/{x['confidence']}"] for x in c["intra_constraints"]]
        o.append(longtable("p{18mm}p{28mm}p{28mm}p{62mm}p{16mm}",
                           ["类型", "先", "后", "理由", "来源/置信"], rows, "工序内硬约束"))

    o.append("\\paragraph{二、零件清单}")
    if c["parts"]:
        rows = [[str(i + 1), p["name"], str(p["qty"])] for i, p in enumerate(c["parts"])]
        o.append(longtable("rp{100mm}r", ["序", "零件名（手册原文）", "数量"], rows))
    else:
        o.append("\\begingroup\\zihao{6}\\noindent 本步无非紧固件零件行。\\par\\endgroup")

    o.append("\\paragraph{三、紧固件清单与推荐扭矩（工具与扭矩均为推断/参考值）}")
    if c["fasteners"]:
        rows = [[f["name"], str(f["qty"]), f["thread"], f["head"], f["tool"], f["torque_class"], f["torque_text"]]
                for f in c["fasteners"]]
        o.append(longtable("p{28mm}rp{10mm}p{28mm}p{34mm}p{8mm}p{22mm}",
                           ["规格（原文）", "数量", "螺纹", "头型（推断）", "工具（推断）", "类", "推荐扭矩"], rows))
        tot = sum((f["qty"] or 0) for f in c["fasteners"])
        o.append("\\begingroup\\zihao{6}\\noindent " + tex(f"本步紧固件合计 {tot} 件。扭矩类别与取值规则见附录 C。")
                 + "\\par\\endgroup")
        if c.get("resin_note"):
            o.append("\\begingroup\\zihao{6}\\noindent " + texb(c["resin_note"]) + "\\par\\endgroup")
    else:
        o.append("\\begingroup\\zihao{6}\\noindent 本步 Workbook 未列紧固件行。\\par\\endgroup")
    tools = sorted({f["tool"] for f in c["fasteners"] if f["tool"] != "无（装配辅件）"})
    if c["official_tools"]:
        tools = [t + "（官方）" for t in c["official_tools"]] + tools
    if tools:
        o.append("\\begingroup\\zihao{6}\\noindent 工具：" + texb("；".join(tools)) + "\\par\\endgroup")
    if c["how_tos"]:
        hts = []
        for h in c["how_tos"]:
            rec = HOWTO.get(h)
            hts.append(f"{h} {rec['title']}：" + " ".join(rec["remarks"])[:300] if rec else h)
        o.append(itemize(hts))

    o.append("\\paragraph{四、装后调整与检验点}")
    items = [a["text"] + "（" + a["src"] + "）" for a in c["adjustments"]]
    for r in c["design_rules"]:
        items.append(f"设计规则：{r['rule']}（依据 {r['basis']}，置信 {r['confidence']}，来自功能线 {r['track']}·{r['unit']}）")
    o.append(itemize(items) if items
             else "\\begingroup\\zihao{6}\\noindent 功能线未给出本步的装后调整项；官方手册无验收字段。\\par\\endgroup")

    o.append("\\paragraph{五、前置条件（本步之前必须完成）}")
    if c["prerequisites"] or c["doc_predecessors"]:
        rows = [[p["op"], p["title"], p["kind"], p["why"], f"{p['source']}/{p['confidence']}"] for p in c["prerequisites"]]
        for p in c["doc_predecessors"]:
            if all(p["op"] != x["op"] for x in c["prerequisites"]):
                rows.append([p["op"], p["title"], "手册序", p["why"], "官方资料/高"])
        o.append(longtable("p{13mm}p{40mm}p{24mm}p{60mm}p{16mm}",
                           ["前置步", "前置步标题", "约束", "理由", "来源/置信"], rows))
    else:
        o.append("\\begingroup\\zihao{6}\\noindent 无工序间前置约束记录（可作为起始步或台架首步）。\\par\\endgroup")
    if c["hints"]:
        rows = [[h["kind"], h["frm"], h["to"], h["why"], f"{h['source']}/{h['confidence']}"] for h in c["hints"]]
        o.append(longtable("p{22mm}p{13mm}p{13mm}p{78mm}p{16mm}",
                           ["提示类型", "from", "to", "说明", "来源/置信"], rows, "顺序提示（安全 / 同台架 / 后序）"))

    o.append("\\paragraph{六、官方 Notes / Remarks 原文}")
    on = []
    if c["official_notes"]:
        on.append("Notes：" + c["official_notes"])
    for r in c["official_remarks"]:
        on.append("Remark：" + r)
    for p in c["official_problems"]:
        on.append("Problems & questions：" + p)
    o.append(itemize(on) if on
             else "\\begingroup\\zihao{6}\\noindent 官方 Workbook 本步的 Notes / Remarks / Problems 栏均为空。\\par\\endgroup")
    if c["function"] or c["principle"]:
        o.append("\\paragraph{七、功能与原理（功能属性手册摘录）}")
        o.append(itemize((c["function"] + c["principle"])[:4]))

    if c["open_questions"] or c["count_mismatch"]:
        o.append("\\paragraph{八、待确认问题}")
        items = [q["text"] + "（" + q["src"] + "）" for q in c["open_questions"]]
        if c["count_mismatch"]:
            items.append("数量对账：" + "；".join(
                f"{m['name']} 手册 {m['sop_qty']} 件 / GLB {m['glb']} / 已绑定 STEP {m['bound']}（{m['status']}）"
                for m in c["count_mismatch"][:6]) + ("…" if len(c["count_mismatch"]) > 6 else "")
                + "（本体 S3 计数对账）")
        o.append(itemize(items))
    return "\n".join(o)


def sec_cards():
    o = ["\\clearpage", "\\section{工艺卡（114 道工序）}"]
    for c in CARDS:
        o.append(card_tex(c))
    return "\n".join(o)


def sec_appendix():
    o = ["\\clearpage", "\\section{附录 A\\quad 紧固件总表}"]
    o.append(tex(f"全机紧固件共 {sum(FAST_TOTAL.values())} 件、{len(FAST_TOTAL)} 个规格，分布在 "
                 f"{sum(1 for c in CARDS if c['fasteners'])} 道工序上。按家族汇总如下。"))
    rows = [[r["family"], str(r["qty"]), r["head"], r["standard"], r["drive"], r["source"]] for r in FAMILY_SUMMARY]
    rows.append(["合计", str(sum(FAST_TOTAL.values())), "", "", "", ""])
    o.append(longtable("p{26mm}rp{28mm}p{56mm}p{12mm}p{14mm}",
                       ["家族", "总数", "头型（推断）", "标准号（推断）", "驱动", "来源"], rows, "紧固件家族汇总"))
    rows = [[r["name"], str(r["qty"]), r["thread"], r["length"], r["tool"],
             "/".join(r["torque_classes"]), "；".join(r["torque_values"]), str(r["n_ops"])] for r in FASTENER_SUMMARY]
    o.append(longtable("p{30mm}rp{9mm}p{12mm}p{34mm}p{8mm}p{28mm}r",
                       ["规格（原文）", "总数", "螺纹", "长度", "工具（推断）", "类", "推荐扭矩（参考值）", "工序数"],
                       rows, "紧固件规格 × 总数 × 工具 × 推荐扭矩"))

    o.append("\\section{附录 B\\quad 工具总表}")
    rows = [[r["name"], str(r["n_ops"]), r["basis"]] for r in TOOL_SUMMARY]
    o.append(longtable("p{46mm}rp{92mm}", ["工具", "涉及工序数", "依据"], rows, "工具总表（按涉及工序数排序）"))
    o.append(itemize([
        "官方 Workbook 只在 WB-01.1 填写了工具（Allen Key 6、Wrench 14），其余 113 步的工具栏全为空，"
        "本表其余项目全部由紧固件头型的标准对边表推断。",
        "两项官方值与标准表存在一档差异：ISO 7380 的 M8 半圆头内六角应为 5 mm、DIN 985 的 M8 尼龙锁母对边应为 13 mm，"
        "官方却写 Allen Key 6 与 Wrench 14。需向设计方确认 B-screw 的实际头型标准与 Lock Nut 的实际对边——"
        "若 B-screw 实为 DIN 912 圆柱头，则全机 695 根 B-screw 的内六角规格需整体上调一档。",
    ], small=False))

    o.append("\\section{附录 C\\quad 扭矩参考表（含依据）}")
    o.append("\\paragraph{依据}" + tex(DOC["torque_reference"]["basis"]))
    o.append("\\paragraph{声明}" + tex("官方 Workbook 的 field_semantics 明确 torque = UNKNOWN，全表无扭矩字段、"
                                     "无螺纹胶要求（How-To H4 的内部备注写道 “Locktite should be mentioned or shown”，"
                                     "说明官方自己也尚未定稿）。下表所有数值均为参考值，不是官方值，投产前必须由设计方确认。"))
    rows = [[t["thread"], t["A"], t["B"], t["C"], t["prevail"], t["F"]] for t in TORQUE_TABLE]
    o.append(longtable("p{16mm}p{26mm}p{28mm}p{28mm}p{24mm}p{26mm}",
                       ["螺纹", "A 钢制螺纹副", f"B 夹铝件 ×{K_ALU:.2f}", f"C 夹树脂件 ×{K_RESIN:.2f}",
                        "锁母自锁力矩", "F 锁紧螺母合计"], rows, "推荐拧紧扭矩参考表（N·m）"))
    rows = [[k, v["name"], v["rule"]] for k, v in TORQUE_CLASS.items()]
    o.append(longtable("p{8mm}p{54mm}p{88mm}", ["类", "适用场合", "取值规则"], rows, "扭矩类别与取值规则"))
    o.append(itemize([
        "A 类基准值为 ISO 898-1 8.8 级、总摩擦系数 μ = 0.14、90% 屈服利用率的常用查表值。"
        "若实际使用 A2-70 不锈钢紧固件，应在此基础上再下调约 15%（M8 约 21 N·m）。",
        DOC["torque_reference"]["cross_check"],
        "B 类适用于本机机架：CNC 铣孔铝方管（壁厚 5 mm）+ 6 mm 铝连接板 + HF-screw M12-20 端面对接，"
        "以及铝角件 + C-screw M12-20/30。本机不使用 T 型螺母（全机 0 件），HowTo H1 对本机不适用。",
        "C 类适用于夹紧光固化树脂 3D 打印件（约 40 种、60 件，含丝杠端轴承座、电机座、肩部接头等半受力件）；"
        "功能线 H 的设计规则要求“树脂接头的拧紧扭矩必须低于同规格金属件，建议加大垫圈分散压强”。",
        "D 类热塑自攻螺钉一次成型，不可反复拆装同一孔。E 类顶丝必须顶在电机轴的平面（D 面）上并加螺纹胶。"
        "F 类尼龙锁紧螺母为一次性件，尼龙侧必须背向螺钉头（How-To H2）。",
        "本表未覆盖：Standoff FF 8mm（无螺纹规格记录）、Set Screw（无规格记录），"
        "以及所有非螺纹连接（压入式窗框角件、磁铁粘接、扎带、快插气管）。",
        "使用方式：先按被连接件中最弱的材料选类别，再按螺纹查值；分次拧紧（30% → 60% → 100%），"
        "对称交叉顺序；长梁与大平板从中间向两端。以上为通用装配惯例（推断），官方无规定。",
    ], small=False))

    o.append("\\section{附录 D\\quad 台架并行计划}")
    pct = sum(len(v) for v in BENCH_OPS.values()) / 114 * 100
    o.append(tex(f"{len(BENCH_PLAN)} 个台架预装组共 {sum(len(v) for v in BENCH_OPS.values())} 步，占 114 步的 {pct:.0f}%。"
                 "分组来自手册源码 Model.jsx 的 preparing 列表（受控旁证），上机结合点来自 S3 的 unit_joins。"
                 "“最早可开工”是该组全部外部前置完成之后的第一步；“最早交付节点”是该组的上机结合步。"))
    rows = [[g["group"], "、".join(g["ops"]), g["join_op"], g["earliest_start_after"],
             "、".join(g["external_prereqs"]) or "无", str(g["part_rows"]), str(g["part_qty"]), g["title"]]
            for g in BENCH_PLAN]
    o.append(longtable("p{17mm}p{38mm}p{15mm}p{18mm}p{26mm}rrp{38mm}",
                       ["组", "工序", "上机步", "最早开工", "外部前置", "行", "件", "内容"], rows, "台架组明细"))
    rows = [[str(w["wave"]), w["gate"], str(w["n_groups"]), "、".join(w["groups"]), w["must_finish_by"]] for w in WAVES]
    o.append(longtable("rp{26mm}rp{60mm}p{24mm}",
                       ["波次", "开工条件", "组数", "台架组", "最早交付节点"], rows, "并行波次"))
    late = [g for g in BENCH_PLAN if g.get("late_prereqs")]
    if late:
        rows = [[g["group"], "、".join(g["late_prereqs"]), g["join_op"]] for g in late]
        o.append(longtable("p{16mm}p{60mm}p{22mm}", ["组", "晚于上机步的前置（疑点）", "上机步"], rows,
                           "疑点：功能线记录的前置步排在本组上机步之后，已从并行窗口计算中剔除，需人工判定"))

    o.append("\\section{附录 E\\quad 工序间顺序约束总表}")
    inter = [c for c in CONSTRAINT_TABLE if c["scope"] == "工序间"]
    o.append(tex(f"共 {len(CONSTRAINTS)} 条顺序约束，其中工序间 {len(inter)} 条、工序内 {len(CONSTRAINTS)-len(inter)} 条；"
                 f"工艺卡的“前置条件”栏共引用 {PRE_COUNT} 条（MUST_BEFORE / ACCESS / WIRE_BEFORE_CLOSE）。"))
    kc = Counter(c["kind"] for c in CONSTRAINT_TABLE)
    o.append(longtable("p{34mm}rp{104mm}", ["类型", "条数", "含义"], [
        ["MUST_BEFORE", str(kc["MUST_BEFORE"]), "from 必须先于 to：几何落座、夹持关系、装配可达性"],
        ["MUST_AFTER", str(kc["MUST_AFTER"]), "后序关系；各功能线书写方向不一致，本规程只列为提示，不参与前置推导"],
        ["ACCESS", str(kc["ACCESS"]), "可达性：一旦 to 完成，from 涉及的区域不可再检修/调整"],
        ["WIRE_BEFORE_CLOSE", str(kc["WIRE_BEFORE_CLOSE"]), "封盖前必须完成的穿线/过线口/线槽作业"],
        ["ADJUST_AFTER", str(kc["ADJUST_AFTER"]), "在 from 之后、to 之前必须完成的调整（找平、对中、张紧、预紧）"],
        ["SAME_BENCH", str(kc["SAME_BENCH"]), "同一台架预装组，应整体成组后再上机"],
        ["SAFETY", str(kc["SAFETY"]), "安全 / 防损：滑块防脱、带压件、气路调压、联锁"],
    ], "约束类型分布"))
    rows = [[c["track"], c["kind"], c["frm"], c["to"], c["why"], f"{c['source']}/{c['confidence']}"] for c in inter]
    o.append(longtable("p{7mm}p{24mm}p{13mm}p{13mm}p{78mm}p{16mm}",
                       ["线", "类型", "from", "to", "理由", "来源/置信"], rows, "工序间顺序约束逐条"))

    o.append("\\section{附录 F\\quad 装后调整与检验点总表}")
    rows = [[a["op"], a["text"], a["src"]] for a in ADJUSTMENTS]
    o.append(longtable("p{13mm}p{96mm}p{40mm}", ["工序", "调整/检验内容", "来源"], rows,
                       f"装后调整与检验点共 {len(ADJUSTMENTS)} 条"))

    o.append("\\section{附录 G\\quad 待确认问题总表}")
    rows = [[q["op"], q["text"], q["src"]] for q in QUESTIONS]
    o.append(longtable("p{13mm}p{100mm}p{36mm}", ["工序", "问题", "来源"], rows, f"待确认问题共 {len(QUESTIONS)} 条"))

    o.append("\\section{附录 H\\quad How-To 原文}")
    used = sorted({h for c in CARDS for h in c["how_tos"]})
    o.append(tex("Workbook 与功能线共引用 " + "、".join(used) + "。以下为 HowTo_Blad1.csv 原文（未改写）。"))
    rows = [[h["id"], h["title"], "；".join(h["tools"]), " ".join(h["remarks"])[:600]]
            for h in HOWTO.values() if h["id"] in used]
    o.append(longtable("p{12mm}p{40mm}p{30mm}p{68mm}", ["编号", "标题", "工具", "要点原文"], rows, "本机用到的 How-To"))
    return "\n".join(o)



TEX = "\n".join([PREAMBLE, COVER, sec_intro(), sec_route(), sec_cards(), sec_appendix(), "\\end{document}"])
OUT_TEX.write_text(TEX, encoding="utf-8")

if "--no-pdf" not in sys.argv:
    ok = True
    for i in (1, 2):
        r = subprocess.run(["xelatex", "-interaction=nonstopmode", "-halt-on-error", "sop.tex"],
                           cwd=HERE, capture_output=True, text=True)
        log = (HERE / "sop.log").read_text(encoding="utf-8", errors="replace") if (HERE / "sop.log").exists() else r.stdout
        errs = [l for l in log.splitlines() if l.startswith("!")]
        print(f"    xelatex pass {i}: rc={r.returncode}  errors={len(errs)}")
        for e in errs[:10]:
            print("      ", e)
        if errs or r.returncode != 0:
            ok = False
            (HERE / f"xelatex{i}.log").write_text(log, encoding="utf-8")
            break
    if ok and (HERE / "sop.pdf").exists():
        (HERE / "sop.pdf").replace(OUT_PDF)
        for ext in (".aux", ".out", ".toc"):
            p = HERE / ("sop" + ext)
            if p.exists():
                p.unlink()
        print(f"[3/3] pdf       {OUT_PDF.stat().st_size/1024:.0f} KB")
    else:
        print("[3/3] pdf       FAILED（见 xelatex*.log）")
        sys.exit(1)
