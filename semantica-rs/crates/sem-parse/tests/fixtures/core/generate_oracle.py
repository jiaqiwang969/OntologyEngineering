"""Generate tests/fixtures/core/oracle.json (and the input files under files/) from the original.

Covers the modules ported directly in sem-parse: config, registry, json_parser, csv_parser,
structured_data_parser, document_parser, mcp_parser, methods and the package exports. Run from
var/semantica-original/oracle/src:

    env -i HOME=../home PATH=../venv/bin:/usr/bin:/bin ../venv/bin/python \
        <crate>/tests/fixtures/core/generate_oracle.py <crate>/tests/fixtures/core

Absolute paths of the fixture directory are written as "$FILES". No network, no credentials.
"""

import email.message  # noqa: F401  (semantica.parse needs it imported, as under pytest)
import json
import os
import sys

sys.path.insert(0, os.getcwd())

OUT_DIR = os.path.abspath(sys.argv[1])
FILES = os.path.join(OUT_DIR, "files")
os.makedirs(FILES, exist_ok=True)

from semantica.utils.progress_tracker import get_progress_tracker  # noqa: E402

get_progress_tracker().displays = []

import semantica.parse as P  # noqa: E402
from semantica.parse import (  # noqa: E402
    CSVParser,
    DocumentParser,
    JSONParser,
    MCPParser,
    MethodRegistry,
    ParseConfig,
    StructuredDataParser,
    list_available_methods,
    method_registry,
)
from semantica.parse import methods as M  # noqa: E402

OUT = {}


def f(name):
    return os.path.join(FILES, name)


def write(name, content, mode="w", **kw):
    with open(f(name), mode, **kw) as fh:
        fh.write(content)
    return f(name)


def norm(value):
    """JSON form of a result with the fixture path replaced."""
    import dataclasses

    def default(o):
        if dataclasses.is_dataclass(o):
            return dataclasses.asdict(o)
        if isinstance(o, bytes):
            return {"__bytes__": list(o)}
        if isinstance(o, (set, tuple)):
            return list(o)
        return repr(o)

    text = json.dumps(value, default=default, ensure_ascii=False)
    text = text.replace(FILES, "$FILES")
    return json.loads(text)


def run(fn, *a, **k):
    try:
        return {"ok": norm(fn(*a, **k))}
    except Exception as e:  # noqa: BLE001
        msg = str(e).replace(FILES, "$FILES")
        return {"error": type(e).__name__, "message": msg}


# ----------------------------------------------------------------------------- inputs
write("data.json", json.dumps({"key": "value", "list": [1, 2, 3]}))
write("nested.json", json.dumps({"user": {"name": "Ann", "tags": ["a", "b"], "address": {"city": "X"}},
                                 "items": [{"name": "i1", "n": 1}, {"name": "i2", "n": 2.5}], "flag": True,
                                 "none": None}))
write("array1.json", json.dumps([{"only": 1}]))
write("array2.json", json.dumps([{"x": 1}, {"x": 2}]))
write("ints.json", json.dumps([1, 2, 3]))
write("scalar.json", json.dumps("text"))
write("invalid.json", '{"a": 1,}')
write("bom.json", '﻿{"a": 1}', encoding="utf-8")
write("latin1.json", '{"name": "caf\xe9"}', encoding="latin-1")
write("badutf8.json", b'{"name": "caf\xe9"}', mode="wb")
write("crlf.json", '{\r\n  "a": [1,\r\n 2]\r\n}\r\n', newline="")
write("companies.csv", "name,company,role\nTim Cook,Apple Inc.,CEO\nSatya Nadella,Microsoft Corporation,CEO\n")
write("semicolon.csv", "a;b;c\n1;2;3\n4;5;6\n")
write("ragged.csv", "h1,h2\n1,2,3,4\n5\n\n6,7\n")
write("quoted.csv", 'name,note\n"Smith, J","said ""hi""\non two lines"\nplain,x\r\n')
write("empty.csv", "")
write("crlf.csv", "a,b\r\n1,2\r\n3,4\r\n", newline="")
write("latin1.csv", "n\ncaf\xe9\n", encoding="latin-1")
write("mixedbytes.csv", b"n,v\nab\xffc,1\n", mode="wb")
write("single.csv", "word\nanother\n")
write("data.yaml", "parse:\n  default_encoding: latin-1\n  extract_tables: false\nparse_methods:\n  document:\n    extract_images: true\nlist:\n  - 1\n  - two\n  - 3.5\nnested:\n  a: null\n  b: [x, y]\n")
write("scalar.yaml", "just a string\n")
write("config.json", json.dumps({"parse": {"ocr_language": "deu", "max": 5}, "parse_methods": {"web": {"render": False}}}))
write("config.toml", '[parse]\ndefault_encoding = "utf-16"\nlimit = 10\n\n[parse_methods.code]\nlanguage = "python"\n')
write("bad.yaml", "parse: [unclosed\n")
write("list.yaml", "- a\n- b\n")
write("notes.txt", "Apple Inc. is a technology company. Tim Cook is the CEO.")
write("crlf.txt", "line1\r\nline2\rline3\n", newline="")
write("latin1.txt", "caf\xe9", encoding="latin-1")
write("unknown.xyz", "whatever")
write("data.xml", '<?xml version="1.0"?>\n<root><person id="1"><name>Alice</name></person></root>\n')

# --------------------------------------------------------------------------- json
jp = JSONParser()
OUT["json"] = {
    "parse_file": run(jp.parse, f("data.json")),
    "parse_nested": run(jp.parse, f("nested.json")),
    "parse_flatten": run(jp.parse, f("nested.json"), flatten=True),
    "parse_paths": run(jp.parse, f("nested.json"), extract_paths=True),
    "parse_flatten_paths": run(jp.parse, f("nested.json"), flatten=True, extract_paths=True),
    "parse_string": run(jp.parse, '{"a": [1, {"b": 2}], "c": 1.0, "d": 1e5}'),
    "parse_array_string": run(jp.parse, '[{"x":1},{"x":2}]'),
    "parse_str_scalar": run(jp.parse, '"hello"'),
    "parse_int_scalar": run(jp.parse, "42"),
    "parse_bool_scalar": run(jp.parse, "true"),
    "parse_null_scalar": run(jp.parse, "null"),
    "parse_float_scalar": run(jp.parse, "1.5"),
    "invalid_string": run(jp.parse, '{"a": }'),
    "invalid_string2": run(jp.parse, "[1, 2,]"),
    "invalid_string3": run(jp.parse, "{\"a\": 1}\nextra"),
    "invalid_file": run(jp.parse, f("invalid.json")),
    "missing_path": run(jp.parse, __import__("pathlib").Path(f("missing.json"))),
    "missing_str": run(jp.parse, f("missing.json")),
    "empty_string": run(jp.parse, ""),
    "long_string": run(jp.parse, json.dumps({"k": "v" * 2000})),
    "bom_file": run(jp.parse, f("bom.json")),
    "latin1_file": run(jp.parse, f("latin1.json"), encoding="latin-1"),
    "badutf8_file": run(jp.parse, f("badutf8.json")),
    "unknown_encoding": run(jp.parse, f("data.json"), encoding="nope"),
    "crlf_file": run(jp.parse, f("crlf.json")),
    "to_dict_object": run(jp.parse_to_dict, f("data.json")),
    "to_dict_array1": run(jp.parse_to_dict, f("array1.json")),
    "to_dict_array2": run(jp.parse_to_dict, f("array2.json")),
    "to_dict_ints": run(jp.parse_to_dict, f("ints.json")),
    "to_dict_scalar": run(jp.parse_to_dict, f("scalar.json")),
    "to_dict_flat_array": run(jp.parse_to_dict, f("array2.json"), flatten=True),
    "to_list_array": run(jp.parse_to_list, f("ints.json")),
    "to_list_object": run(jp.parse_to_list, f("data.json")),
    "to_list_scalar": run(jp.parse_to_list, f("scalar.json")),
    "to_list_flat_array": run(jp.parse_to_list, f("array2.json"), flatten=True),
    "values_user_name": run(jp.extract_values, f("nested.json"), "user.name"),
    "values_items_name": run(jp.extract_values, f("nested.json"), "items.name"),
    "values_missing": run(jp.extract_values, f("nested.json"), "user.missing"),
    "values_tags": run(jp.extract_values, f("nested.json"), "user.tags"),
    "values_empty_path": run(jp.extract_values, '{"": 1}', ""),
}

# ---------------------------------------------------------------------------- csv
cp = CSVParser()
OUT["csv"] = {
    "basic": run(cp.parse, f("companies.csv")),
    "no_header": run(cp.parse, f("companies.csv"), has_header=False),
    "skip_max": run(cp.parse, f("companies.csv"), skip_rows=1, max_rows=1),
    "no_header_skip": run(cp.parse, f("companies.csv"), has_header=False, skip_rows=1),
    "max_zero": run(cp.parse, f("companies.csv"), max_rows=0),
    "max_negative": run(cp.parse, f("companies.csv"), max_rows=-1),
    "semicolon": run(cp.parse, f("semicolon.csv"), delimiter=";"),
    "sniff": run(cp.parse, f("semicolon.csv"), delimiter=None),
    "sniff_fail": run(cp.parse, f("ragged.csv"), delimiter=None),
    "sniff_single": run(cp.parse, f("single.csv"), delimiter=None),
    "ragged": run(cp.parse, f("ragged.csv")),
    "ragged_no_header": run(cp.parse, f("ragged.csv"), has_header=False),
    "quoted": run(cp.parse, f("quoted.csv")),
    "empty": run(cp.parse, f("empty.csv")),
    "empty_no_header": run(cp.parse, f("empty.csv"), has_header=False),
    "crlf": run(cp.parse, f("crlf.csv")),
    "latin1": run(cp.parse, f("latin1.csv"), encoding="latin-1"),
    "mixedbytes": run(cp.parse, f("mixedbytes.csv")),
    "missing": run(cp.parse, f("missing.csv")),
    "bad_delimiter": run(cp.parse, f("companies.csv"), delimiter=";;"),
    "unknown_encoding": run(cp.parse, f("companies.csv"), encoding="nope"),
    "has_header_zero": run(cp.parse, f("companies.csv"), has_header=0),
    "to_dict": run(cp.parse_to_dict, f("ragged.csv")),
    "to_dict_delim": run(cp.parse_to_dict, f("semicolon.csv"), delimiter=";"),
    "to_list": run(cp.parse_to_list, f("ragged.csv")),
    "to_list_no_header": run(cp.parse_to_list, f("ragged.csv"), has_header=False),
}

# ---------------------------------------------------------------- structured data
sp = StructuredDataParser()
OUT["structured"] = {
    "json_auto": run(sp.parse_data, f("data.json")),
    "json_explicit": run(sp.parse_data, f("data.json"), data_format="json"),
    "csv_auto": run(sp.parse_data, f("companies.csv")),
    "csv_delimiter": run(sp.parse_data, f("semicolon.csv"), data_format="csv", delimiter=";"),
    "yaml_file": run(sp.parse_data, f("data.yaml")),
    "yaml_scalar": run(sp.parse_data, f("scalar.yaml")),
    "yaml_list": run(sp.parse_data, f("list.yaml")),
    "yaml_string": run(sp.parse_data, "a: 1\nb: [x, 2]\n", data_format="yaml"),
    "json_content": run(sp.parse_data, ' {"k": [1, 2]}'),
    "csv_content": run(sp.parse_data, "a,b\n1,2"),
    "xml_content_detect": run(sp._detect_format, "<root/>"),
    "unknown_content": run(sp.parse_data, "plain words"),
    "unknown_file": run(sp.parse_data, f("unknown.xyz")),
    "bad_format": run(sp.parse_data, f("data.json"), data_format=","),
    "validate_none": run(sp.validate_data, {"a": 1}, None),
    "validate_required": run(sp.validate_data, {"a": 1}, {"required": ["a", "b"]}),
    "validate_types": run(sp.validate_data,
                          {"s": "x", "n": 1.5, "i": True, "b": True, "l": [1], "o": {}, "z": None, "q": 1},
                          {"properties": {"s": {"type": "string"}, "n": {"type": "number"}, "i": {"type": "integer"},
                                          "b": {"type": "boolean"}, "l": {"type": "array"}, "o": {"type": "object"},
                                          "z": {"type": "null"}, "q": {}}}),
    "validate_number_int": run(sp.validate_data, {"n": 3}, {"properties": {"n": {"type": "number"}}}),
    "validate_non_dict": run(sp.validate_data, [1, 2], {"required": ["a"]}),
    "convert_dict_json": run(sp.convert_data, {"name": "Ann", "tags": ["a", "é"], "n": 1.0}, "json", "json"),
    "convert_dict_csv": run(sp.convert_data, {"a": 1, "b": "x,y", "c": None}, "json", "csv"),
    "convert_list_dicts_csv": run(sp.convert_data, [{"a": 1, "b": 2}, {"a": 3}], "json", "csv"),
    "convert_list_extra_csv": run(sp.convert_data, [{"a": 1}, {"a": 2, "z": 3}], "json", "csv"),
    "convert_list_lists_csv": run(sp.convert_data, [[1, "two"], ["x\"y", None]], "json", "csv"),
    "convert_list_mixed_csv": run(sp.convert_data, [[1], 5], "json", "csv"),
    "convert_empty_csv": run(sp.convert_data, [], "json", "csv"),
    "convert_scalar_csv": run(sp.convert_data, 42, "json", "csv"),
    "convert_dict_xml": run(sp.convert_data, {"person": {"name": "A & B", "age": 3, "tags": ["x", {"k": "v"}], "empty": [], "none": None}}, "json", "xml"),
    "convert_list_xml": run(sp.convert_data, [1, {"a": "<b>"}], "json", "xml"),
    "convert_scalar_xml": run(sp.convert_data, 3.5, "json", "xml"),
    "convert_bad_tag_xml": run(sp.convert_data, {"first name": 1}, "json", "xml"),
    "convert_empty_dict_xml": run(sp.convert_data, {"a": {}, "b": ""}, "json", "xml"),
    "convert_dict_yaml": run(sp.convert_data, {"b": [1, 2], "a": {"x": "y"}}, "json", "yaml"),
    "convert_file_csv": run(sp.convert_data, f("array2.json"), "json", "csv"),
    "convert_string_yaml": run(sp.convert_data, '{"z": 1, "a": [true, null]}', "json", "yaml"),
    "convert_csv_file_json": run(sp.convert_data, f("companies.csv"), "csv", "json"),
    "convert_unsupported": run(sp.convert_data, {"a": 1}, "json", "toml"),
    "schema_dict": run(sp.extract_schema, {"s": "x", "i": 1, "f": 1.5, "b": False, "n": None, "l": [{"a": 1}], "o": {"p": []}}),
    "schema_list": run(sp.extract_schema, [1, 2]),
    "schema_empty_list": run(sp.extract_schema, []),
    "schema_scalar": run(sp.extract_schema, "x"),
    "schema_none": run(sp.extract_schema, None),
}

# -------------------------------------------------------------------------- config
pc_yaml = ParseConfig(f("data.yaml"))
pc_json = ParseConfig(f("config.json"))
pc_toml = ParseConfig(f("config.toml"))
pc_bad = ParseConfig(f("bad.yaml"))
pc_list = ParseConfig(f("list.yaml"))
pc_missing = ParseConfig(f("nope.yaml"))
saved = dict(os.environ)
os.environ.update({"PARSE_EXTRACT_TABLES": "Yes", "PARSE_OCR_LANGUAGE": "fra", "PARSE_MAX_PAGES": "12",
                   "PARSE_THRESHOLD": "0.75", "PARSE_FLAG": "TRUE", "PARSE_NAME": "custom", "PARSE_EMPTY": "",
                   "PARSE_EXP": "1e3", "PARSE_NEG": "-5", "OTHER": "x"})
pc_env = ParseConfig()
env_get = {
    "configured": pc_env.get("max_pages"),
    "missing_default": pc_env.get("missing", "dflt"),
    "missing_none": pc_env.get("missing"),
}
os.environ["PARSE_TIMEOUT"] = "30"
os.environ["PARSE_RATIO"] = "2.5"
os.environ["PARSE_ON"] = "1"
os.environ["PARSE_WORD"] = "yes"
env_get.update({
    "int_default": pc_env.get("timeout", 5),
    "float_default": pc_env.get("ratio", 1.0),
    "bool_default_digit": pc_env.get("on", False),
    "bool_default_word": pc_env.get("word", False),
    "str_default": pc_env.get("word", "x"),
    "int_default_bad": pc_env.get("word", 7),
})
os.environ.clear()
os.environ.update(saved)
pc_api = ParseConfig()
pc_api.set("default_encoding", "ascii")
pc_api.set_method_config("document", extract_tables=False, pages=[1, 2])
OUT["config"] = {
    "yaml": norm(pc_yaml.get_all()),
    "json": norm(pc_json.get_all()),
    "toml": norm(pc_toml.get_all()),
    "bad_yaml": norm(pc_bad.get_all()),
    "list_yaml": norm(pc_list.get_all()),
    "missing": norm(pc_missing.get_all()),
    "env": norm({k: v for k, v in pc_env.get_all()["config"].items()}),
    "env_get": norm(env_get),
    "api_get": pc_api.get("default_encoding"),
    "api_method": norm(pc_api.get_method_config("document")),
    "api_method_missing": norm(pc_api.get_method_config("web")),
    "api_all": norm(pc_api.get_all()),
    "yaml_get_method": norm(pc_yaml.get_method_config("document")),
}

# ------------------------------------------------------------------------ registry
reg_log = {"initial": norm(list_available_methods()), "task": norm(list_available_methods("document")),
           "unknown_task": norm(list_available_methods("nope")), "empty_task": norm(list_available_methods(""))}
other = MethodRegistry()
other.register("document", "mine", lambda *a, **k: {"mine": True})
other.register("newtask", "m1", lambda *a, **k: None)
reg_log["after_register"] = norm(method_registry.list_all())
reg_log["get_mine"] = method_registry.get("document", "mine") is not None
reg_log["get_missing"] = method_registry.get("document", "zzz") is None
reg_log["get_missing_task"] = method_registry.get("zzz", "default") is None
method_registry.unregister("document", "mine")
method_registry.unregister("document", "never")
method_registry.unregister("never", "x")
reg_log["after_unregister"] = norm(method_registry.list_all())
saved_methods = {t: dict(m) for t, m in MethodRegistry._methods.items()}
method_registry.clear("web")
reg_log["after_clear_web"] = norm(method_registry.list_all())
method_registry.clear()
reg_log["after_clear_all"] = norm(method_registry.list_all())
MethodRegistry._methods.clear()
MethodRegistry._methods.update(saved_methods)
MethodRegistry._methods.pop("newtask", None)
reg_log["restored"] = norm(method_registry.list_all())
OUT["registry"] = reg_log

# ------------------------------------------------------------------------ document
dp = DocumentParser()
OUT["document"] = {
    "parse_txt": run(dp.parse_document, f("notes.txt")),
    "parse_txt_type": run(dp.parse_document, f("unknown.xyz"), file_type="text"),
    "parse_alias": run(dp.parse, f("notes.txt")),
    "extract_text": run(dp.extract_text, f("notes.txt")),
    "extract_text_crlf": run(dp.extract_text, f("crlf.txt")),
    "extract_text_latin1": run(dp.extract_text, f("latin1.txt"), encoding="latin-1"),
    "extract_text_utf8_ignore": run(dp.extract_text, f("latin1.txt")),
    "extract_text_missing": run(dp.extract_text, f("missing.txt")),
    "extract_text_unknown": run(dp.extract_text, f("unknown.xyz")),
    "extract_metadata": run(dp.extract_metadata, f("notes.txt")),
    "unknown_format": run(dp.parse_document, f("unknown.xyz")),
    "missing": run(dp.parse_document, f("missing.txt")),
    "bad_type": run(dp.parse_document, f("notes.txt"), file_type="rtf"),
    "batch": run(dp.parse_batch, [f("notes.txt"), f("missing.txt"), f("unknown.xyz"), f("crlf.txt")]),
    "batch_stop": run(dp.parse_batch, [f("notes.txt"), f("missing.txt"), f("crlf.txt")], continue_on_error=False),
    "batch_empty": run(dp.parse_batch, []),
    "parse_list": run(dp.parse, [f("notes.txt"), f("missing.txt"), f("crlf.txt")]),
    "parse_list_bad_item": run(dp.parse, [f("notes.txt"), 5]),
    "parse_encoding": run(dp.parse_document, f("latin1.txt"), encoding="latin-1"),
    "supported_formats": dp.supported_formats,
}

# ------------------------------------------------------------------------------ mcp
mp = MCPParser()
OUT["mcp"] = {
    "resource_json": run(mp.parse_resource_response, {"uri": "file:///x", "mimeType": "application/json",
                                                       "contents": [{"type": "text", "text": ' {"a": [1, 2]}'}],
                                                       "metadata": {"k": 1}}),
    "resource_multi": run(mp.parse_resource_response, {"contents": [
        {"type": "text", "text": "name,age\nAnn,3\n\nBob,4,extra\nCy,5"},
        {"type": "text", "text": "a\tb\n1\t2"},
        {"type": "text", "text": "single line, with comma"},
        {"type": "text", "text": "[1, 2"},
        {"type": "image", "data": "aGVsbG8=", "mimeType": "image/png"},
        {"type": "image", "data": ""},
        {"type": "resource", "resource": {"uri": "r"}},
        {"type": "blob", "data": [1, 2]},
        {"text": "default type"},
        "not a dict",
        {"type": "text", "text": 5},
    ]}),
    "resource_empty": run(mp.parse_resource_response, {"contents": []}),
    "resource_none": run(mp.parse_resource_response, {"contents": None, "uri": "u"}),
    "resource_not_dict": run(mp.parse_resource_response, [1, 2]),
    "resource_int_contents": run(mp.parse_resource_response, {"contents": 5}),
    "resource_str_contents": run(mp.parse_resource_response, {"contents": "ab"}),
    "tool_output": run(mp.parse_tool_output, {"content": [{"type": "text", "text": "{\"ok\": true}"}], "isError": True}),
    "tool_output_default": run(mp.parse_tool_output, {}),
    "mcp_resource": run(mp.parse_mcp_data, {"contents": [{"type": "text", "text": "hi"}]}),
    "mcp_tool": run(mp.parse_mcp_data, {"content": []}),
    "mcp_other_dict": run(mp.parse_mcp_data, {"x": 1}),
    "mcp_list": run(mp.parse_mcp_data, [1]),
    "mcp_unknown_type": run(mp.parse_mcp_data, {"x": 1}, data_type="weird"),
    "mcp_explicit_resource": run(mp.parse_mcp_data, {"contents": [{"text": "x"}]}, data_type="resource"),
    "mcp_resource_error": run(mp.parse_mcp_data, "str", data_type="resource"),
}

# -------------------------------------------------------------------------- methods
OUT["methods"] = {
    "list_available": norm(M.list_available_methods()),
    "parse_document_txt": run(M.parse_document, f("notes.txt")),
    "parse_document_missing": run(M.parse_document, f("missing.txt")),
    "parse_structured_json": run(M.parse_structured_data, f("data.json")),
    "parse_structured_bad": run(M.parse_structured_data, f("unknown.xyz")),
    "parse_json_via_structured": run(M.parse_json, f("data.json")),
    "parse_json_string": run(M.parse_json, '{"a": 1}'),
    "parse_csv_fallback": run(M.parse_csv, f("companies.csv")),
    "parse_csv_semicolon": run(M.parse_csv, f("semicolon.csv"), delimiter=";"),
    "parse_media_attr": run(M.parse_media, f("notes.txt")),
    "parse_pdf_txt": run(M.parse_pdf, f("notes.txt")),
    "parse_docx_txt": run(M.parse_docx, f("notes.txt")),
    "get_parse_method_default": M.get_parse_method("document", "default") is not None,
    "get_parse_method_missing": M.get_parse_method("document", "zzz") is None,
}
calls = []


def custom_ok(*args, **kwargs):
    calls.append({"args": [str(a) for a in args], "kwargs": kwargs})
    return {"custom": True, "nargs": len(args)}


def custom_fail(*args, **kwargs):
    raise ValueError("boom")


method_registry.register("document", "custom_ok", custom_ok)
method_registry.register("document", "custom_fail", custom_fail)
method_registry.register("structured", "custom_ok", custom_ok)
OUT["methods"]["custom_ok_document"] = run(M.parse_document, f("notes.txt"), method="custom_ok", extra=1)
OUT["methods"]["custom_fail_document"] = run(M.parse_document, f("notes.txt"), method="custom_fail")
OUT["methods"]["custom_ok_pdf"] = run(M.parse_pdf, f("notes.txt"), method="custom_ok")
OUT["methods"]["custom_ok_csv"] = run(M.parse_csv, f("companies.csv"), method="custom_ok")
OUT["methods"]["custom_calls"] = norm(calls)
OUT["methods"]["unknown_method"] = run(M.parse_document, f("notes.txt"), method="nope")
method_registry.unregister("document", "custom_ok")
method_registry.unregister("document", "custom_fail")
method_registry.unregister("structured", "custom_ok")

# The method config aliasing: get_method_config returns the stored dict and the
# convenience functions update it in place.
from semantica.parse.config import parse_config  # noqa: E402

parse_config.set_method_config("structured", seen=1)
M.parse_structured_data(f("data.json"), extra_kw=2)
OUT["methods"]["method_config_after"] = norm(parse_config.get_method_config("structured"))
parse_config._method_configs.pop("structured", None)

# ----------------------------------------------------------- cross-module routes
write("page.html", "<html><head><title>Route Page</title><meta name='description' content='d'></head>"
                   "<body><h1 id='top'>Head</h1><p>Hello <a href='/x' rel='nofollow'>link</a></p>"
                   "<table><tr><th>a</th></tr><tr><td>1</td></tr></table><script>var x=1;</script></body></html>")
write("route.xml", '<?xml version="1.0"?>\n<companies><company name="A"/><company name="B">text</company></companies>\n')
write("script.py", 'import os\nfrom sys import path as p\n\n# comment one\ndef f(a, b=1):\n    """doc"""\n    return a\n\nclass C(object):\n    pass\n')
write("mail.eml", "From: a@example.com\nTo: b@example.com, c@example.com\nSubject: Route test\nDate: Mon, 1 Jan 2024 10:00:00 +0000\n\nBody text.\n")
import docx  # noqa: E402

d = docx.Document()
d.core_properties.title = "Route DOCX"
d.add_heading("Heading One", level=1)
d.add_paragraph("First para.")
d.save(f("route.docx"))
from PIL import Image  # noqa: E402

Image.new("RGB", (4, 3), (255, 0, 0)).save(f("tiny.png"))
dp2 = DocumentParser()
OUT["routes"] = {
    "document_html": run(dp2.parse_document, f("page.html")),
    "document_html_text": run(dp2.extract_text, f("page.html")),
    "document_html_text_raw": run(dp2.extract_text, f("page.html"), clean=False),
    "document_html_metadata": run(dp2.extract_metadata, f("page.html")),
    "document_docx": run(dp2.parse_document, f("route.docx")),
    "document_docx_text": run(dp2.extract_text, f("route.docx")),
    "document_docx_metadata": run(dp2.extract_metadata, f("route.docx")),
    "web_html": run(M.parse_web_content, "<p>x <a href='y'>z</a></p>", base_url="http://h/"),
    "web_xml": run(M.parse_web_content, f("route.xml"), content_type="xml"),
    "web_bad_type": run(M.parse_web_content, "<p/>", content_type="json"),
    "xml_via_structured": run(M.parse_xml, f("route.xml")),
    "structured_xml": run(M.parse_structured_data, f("route.xml")),
    "image_via_media": run(M.parse_image, f("tiny.png")),
    "media_attr": run(M.parse_media, f("tiny.png")),
    "code": run(M.parse_code, f("script.py")),
    "email": run(M.parse_email, f("mail.eml")),
    "email_text": run(M.parse_email, "Subject: S\n\nhi\n"),
    "convert_xml_json": run(sp.convert_data, f("route.xml"), "xml", "json"),
}
_cases = [{"first name": 1}, {"1abc": 1}, {"": 1}, {"a b c": 1}, {"ok": {"bad key": 2}}, {"x": "\x01"},
          {"a:b": 1}, {"xml:lang": "en"}, {"é": 1}, {"a/b": 1}, {"a&b": 1}, {"-x": 1}, {".x": 1},
          {"x-1.y_z": 1}, {"a": [{"b c": 1}]}, {"tab\tkey": 1}, {"k": "a\r\nb\rc"}, {"a<b": 1},
          {"a>b": 1}, {"a\"b": 1}, {"aé b": 1}, [{"x y": 1}], {"xmlns": 1}, {"xmlns:a": 1}, {":a": 1},
          {"a:": 1}, {"a:b:c": 1}, {"k": "line1\nline2"}, {"k": True, "n": None, "f": 1.5, "big": 10**15}]
OUT["xml_conversions"] = [{"input": c, "result": run(sp.convert_data, c, "json", "xml")} for c in _cases]

# ------------------------------------------------------------- parse_provenance
from semantica.parse.parse_provenance import ParserWithProvenance  # noqa: E402

OUT["provenance"] = {
    "construct_default": run(ParserWithProvenance),
    "construct_tracking": run(ParserWithProvenance, provenance=True, agent_id="agent-x", is_automated=False, x=1),
}

OUT["exports"] = {"all": list(P.__all__)}

with open(os.path.join(OUT_DIR, "oracle.json"), "w") as fh:
    json.dump(OUT, fh, indent=1, ensure_ascii=False, sort_keys=False)
print("wrote", os.path.join(OUT_DIR, "oracle.json"))
