"""Generate the office fixtures and tests/fixtures/office/oracle.json from the original Semantica.

Creates small .docx/.pptx/.xlsx files (with python-docx, python-pptx, openpyxl and hand-written
packages for edge cases) next to this script, runs semantica.parse's DOCXParser, PPTXParser and
ExcelParser on them and records every result or exception. Run from
var/semantica-original/oracle/src:

    env -i HOME=../home PATH=../venv/bin:/usr/bin:/bin ../venv/bin/python \
        <crate>/tests/fixtures/office/generate_oracle.py

Paths inside messages are replaced by "{DIR}" (this directory) so the expectations do not depend
on where the crate lives; values that depend on the current time are replaced by "{NOW}".
"""

import dataclasses
import datetime as dt
import io
import json
import math
import os
import re
import sys
import zipfile

sys.path.insert(0, os.getcwd())

import email.message  # noqa: E402,F401  (semantica.parse needs it imported first)

from semantica.utils.progress_tracker import get_progress_tracker  # noqa: E402

get_progress_tracker().displays = []

import docx  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_UNDERLINE  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.shared import Pt  # noqa: E402

import pptx  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt as PPt  # noqa: E402

import openpyxl  # noqa: E402

from semantica.parse.docx_parser import DOCXParser  # noqa: E402
from semantica.parse.excel_parser import ExcelParser  # noqa: E402
from semantica.parse.pptx_parser import PPTXParser  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = {"docx": [], "pptx": [], "excel": []}
NOW_FIELDS = set()


def p(name):
    return os.path.join(HERE, name)


def jsonable(value):
    """JSON form of a parser result: dataclasses as dicts, dict keys as str(key),
    datetimes/Timestamps as str(), NaN as the string "NaN"."""
    if dataclasses.is_dataclass(value) and not isinstance(value, type):
        return {f.name: jsonable(getattr(value, f.name)) for f in dataclasses.fields(value)}
    if isinstance(value, dict):
        return {str(k): jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [jsonable(v) for v in value]
    if isinstance(value, bool) or value is None or isinstance(value, str):
        return value
    if isinstance(value, int):
        return int(value)
    if isinstance(value, float):
        if math.isnan(value):
            return "NaN"
        return float(value)
    if hasattr(value, "item") and not hasattr(value, "isoformat"):
        return jsonable(value.item())
    return str(value)


def mask_now(case):
    """Mask the current time python-docx/python-pptx put in default core properties."""
    ok = case.get("ok")
    if isinstance(ok, dict):
        meta = ok.get("metadata")
        if isinstance(meta, dict) and meta.get("last_modified_by") in ("python-docx", "python-pptx"):
            meta["modified"] = "{NOW}"
    return case


def scrub(value):
    """Replace this directory with {DIR} in strings."""
    if isinstance(value, str):
        return value.replace(HERE, "{DIR}")
    if isinstance(value, list):
        return [scrub(v) for v in value]
    if isinstance(value, dict):
        return {scrub(k): scrub(v) for k, v in value.items()}
    return value


def run(section, name, fn, **meta):
    try:
        result = {"ok": jsonable(fn())}
    except Exception as e:  # noqa: BLE001
        result = {"error": type(e).__name__, "message": str(e)}
    case = {"name": name, **meta, **result}
    OUT[section].append(mask_now(scrub(case)))


# ============================================================================ DOCX fixtures


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def make_basic_docx():
    d = docx.Document()
    cp = d.core_properties
    cp.title = "Quarterly Report"
    cp.author = "Alice Analyst"
    cp.subject = "Finance"
    cp.keywords = "q3, revenue"
    cp.comments = "Draft for review"
    cp.category = "Reports"
    cp.last_modified_by = "Bob"
    cp.created = dt.datetime(2024, 1, 2, 3, 4, 5)
    cp.modified = dt.datetime(2024, 2, 3, 4, 5, 6)
    d.add_paragraph("Preamble before any heading.")
    d.add_heading("Executive Summary", level=1)
    para = d.add_paragraph("Revenue grew ")
    r = para.add_run("strongly")
    r.bold = True
    r = para.add_run(" this quarter")
    r.italic = True
    r.underline = True
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para2 = d.add_paragraph()
    r = para2.add_run("Sized text")
    r.font.size = Pt(10.5)
    r.font.name = "Arial"
    r = para2.add_run(" double")
    r.underline = WD_UNDERLINE.DOUBLE
    r = para2.add_run(" none")
    r.underline = False
    r.bold = False
    para2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para3 = d.add_paragraph("Left aligned")
    para3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    d.add_heading("Details", level=2)
    d.add_paragraph("First bullet", style="List Bullet")
    d.add_paragraph("Second bullet", style="List Bullet")
    d.add_paragraph("Numbered", style="List Number")
    tabbed = d.add_paragraph("Name")
    run = tabbed.add_run()
    run.add_tab()
    tabbed.add_run("Value")
    run = tabbed.add_run("line one")
    run.add_break()
    run.add_text("line two")
    run.add_break(WD_BREAK.PAGE)
    run.add_text("after page break")
    linked = d.add_paragraph("See ")
    add_hyperlink(linked, "the website", "https://example.com")
    linked.add_run(" for more.")
    d.add_paragraph("   ")
    d.add_paragraph("")
    t = d.add_table(rows=3, cols=3)
    t.cell(0, 0).text = "Region"
    t.cell(0, 1).text = "Q1"
    t.cell(0, 2).text = "Q2"
    t.cell(1, 0).text = "North"
    t.cell(1, 1).text = " 10 "
    t.cell(1, 2).text = "12"
    t.cell(2, 0).text = "South"
    t.cell(1, 1).merge(t.cell(1, 2))
    t.cell(1, 0).merge(t.cell(2, 0))
    c = t.cell(2, 2)
    c.text = "line a"
    c.add_paragraph("line b")
    d.add_heading("Heading Three", level=3)
    d.add_paragraph("Under three.")
    t2 = d.add_table(rows=1, cols=2)
    t2.cell(0, 0).text = "x"
    d.add_heading("Title Style", level=0)
    d.add_paragraph("Last paragraph.")
    d.save(p("basic.docx"))


W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
R_NS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'

CT_DOCX = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
</Types>"""

PKG_RELS_DOC = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
%s
</Relationships>"""

CORE_REL = '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'

DOC_RELS_STYLES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>"""


def write_zip(name, members):
    with zipfile.ZipFile(p(name), "w", zipfile.ZIP_DEFLATED) as z:
        for member, data in members.items():
            z.writestr(member, data)


def docx_package(name, body, styles=None, core=None, extra=None):
    members = {"[Content_Types].xml": CT_DOCX}
    members["_rels/.rels"] = PKG_RELS_DOC % (CORE_REL if core is not None else "")
    members["word/document.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {W_NS} {R_NS}><w:body>{body}</w:body></w:document>"
    )
    if styles is not None:
        members["word/_rels/document.xml.rels"] = DOC_RELS_STYLES
        members["word/styles.xml"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f"<w:styles {W_NS}>{styles}</w:styles>"
        )
    if core is not None:
        members["docProps/core.xml"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">' + core + "</cp:coreProperties>"
        )
    members.update(extra or {})
    write_zip(name, members)


STYLES_EDGE = """
<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>
<w:style w:type="paragraph" w:styleId="Heading3"><w:name w:val="heading 3"/></w:style>
<w:style w:type="paragraph" w:styleId="Head"><w:name w:val="Heading"/></w:style>
<w:style w:type="paragraph" w:styleId="HeadingX"><w:name w:val="Heading Extra"/></w:style>
<w:style w:type="character" w:styleId="Strong"><w:name w:val="Strong"/></w:style>
<w:style w:type="paragraph" w:styleId="Caption"><w:name w:val="caption"/></w:style>
<w:style w:type="paragraph" w:default="1" w:styleId="Body"><w:name w:val="Body Text"/></w:style>
"""


def run_xml(inner, rpr=""):
    return f"<w:r>{rpr}{inner}</w:r>"


EDGE_BODY = "".join(
    [
        '<w:p><w:pPr><w:pStyle w:val="Heading3"/><w:jc w:val="right"/></w:pPr>'
        + run_xml("<w:t>Edge Heading</w:t>")
        + "</w:p>",
        "<w:p>"
        + run_xml('<w:t xml:space="preserve">a </w:t><w:tab/><w:t>b</w:t><w:br/><w:t>c</w:t>'
                  '<w:br w:type="page"/><w:t>d</w:t><w:br w:type="column"/><w:cr/><w:noBreakHyphen/>'
                  "<w:ptab/><w:t/><w:softHyphen/><w:sym/>")
        + '<w:ins w:id="1"><w:r><w:t>inserted</w:t></w:r></w:ins>'
        + "<w:smartTag><w:r><w:t>smart</w:t></w:r></w:smartTag>"
        + "<w:hyperlink><w:r><w:t>link</w:t></w:r><w:ins><w:r><w:t>x</w:t></w:r></w:ins></w:hyperlink>"
        + run_xml("<w:t>tail</w:t><w:instrText>PAGE</w:instrText><w:delText>gone</w:delText>")
        + "</w:p>",
        '<w:p><w:pPr><w:pStyle w:val="Strong"/><w:jc w:val="both"/></w:pPr>'
        + run_xml("<w:t>char style id falls back to default</w:t>",
                  '<w:rPr><w:b w:val="0"/><w:i w:val="on"/><w:u w:val="words"/><w:sz w:val="21"/>'
                  '<w:rFonts w:ascii=""/></w:rPr>')
        + run_xml("<w:t>pt size</w:t>", '<w:rPr><w:b/><w:u/><w:sz w:val="12pt"/>'
                  '<w:rFonts w:ascii="Courier New" w:hAnsi="X"/></w:rPr>')
        + run_xml("<w:t>zero size</w:t>", '<w:rPr><w:sz w:val="0"/><w:u w:val="none"/></w:rPr>')
        + "</w:p>",
        '<w:p><w:pPr><w:pStyle w:val="Head"/></w:pPr>' + run_xml("<w:t>Plain Heading</w:t>") + "</w:p>",
        '<w:p><w:pPr><w:pStyle w:val="Missing"/><w:jc w:val="center"/></w:pPr>'
        + run_xml("<w:t>unknown style id</w:t>") + "</w:p>",
        '<w:p><w:pPr><w:pStyle w:val="HeadingX"/></w:pPr>' + run_xml("<w:t>Heading Extra Text</w:t>") + "</w:p>",
        '<w:p><w:pPr><w:pStyle w:val="Caption"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr>'
        + run_xml("<w:t>  Caption text  </w:t>") + "</w:p>",
        "<w:tbl><w:tblGrid><w:gridCol/><w:gridCol/><w:gridCol/><w:gridCol/></w:tblGrid>"
        "<w:tr><w:trPr><w:gridBefore w:val=\"1\"/></w:trPr>"
        '<w:tc><w:tcPr><w:gridSpan w:val="2"/><w:vMerge w:val="restart"/></w:tcPr><w:p>' + run_xml("<w:t>top</w:t>") + "</w:p></w:tc>"
        "<w:tc><w:p>" + run_xml("<w:t>r</w:t>") + "</w:p></w:tc></w:tr>"
        "<w:tr><w:trPr><w:gridBefore w:val=\"1\"/></w:trPr>"
        '<w:tc><w:tcPr><w:gridSpan w:val="2"/><w:vMerge/></w:tcPr><w:p/></w:tc>'
        "<w:tc><w:p>" + run_xml("<w:t>s</w:t>") + "</w:p><w:p>" + run_xml("<w:t>t</w:t>") + "</w:p></w:tc></w:tr>"
        "<w:tr><w:tc><w:p>" + run_xml("<w:t>u</w:t>") + "</w:p></w:tc>"
        '<w:tc><w:tcPr><w:gridSpan w:val="0"/></w:tcPr><w:p>' + run_xml("<w:t>zero span</w:t>") + "</w:p></w:tc></w:tr>"
        "</w:tbl>",
        "<w:tbl><w:tblGrid/></w:tbl>",
        '<w:sdt><w:sdtContent><w:p>' + run_xml("<w:t>in sdt</w:t>") + "</w:p></w:sdtContent></w:sdt>",
        "<w:p>" + run_xml("<w:t>Final &amp; done</w:t>") + "</w:p>",
    ]
)

EDGE_CORE = (
    "<dc:title>Edge Doc</dc:title><dc:creator></dc:creator>"
    '<dcterms:created xsi:type="dcterms:W3CDTF">2020-05-06T07:08:09+02:00</dcterms:created>'
    '<dcterms:modified xsi:type="dcterms:W3CDTF">not a date</dcterms:modified>'
    "<cp:lastModifiedBy>Zed</cp:lastModifiedBy><cp:keywords>k</cp:keywords>"
)


def make_docx_edge_cases():
    docx_package("edge.docx", EDGE_BODY, styles=STYLES_EDGE, core=EDGE_CORE)
    # No styles part and no core properties: python-docx falls back to its defaults.
    docx_package(
        "nostyles.docx",
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>' + run_xml("<w:t>Looks like a heading</w:t>") + "</w:p>"
        + "<w:p>" + run_xml("<w:t>Body</w:t>") + "</w:p>",
    )
    # Dates: year only, year-month, empty created element (python-docx raises TypeError).
    docx_package("dates.docx", "<w:p>" + run_xml("<w:t>x</w:t>") + "</w:p>", styles=STYLES_EDGE,
                 core='<dcterms:created xsi:type="dcterms:W3CDTF">2003</dcterms:created>'
                      '<dcterms:modified xsi:type="dcterms:W3CDTF">2003-12</dcterms:modified>')
    docx_package("emptydate.docx", "<w:p>" + run_xml("<w:t>x</w:t>") + "</w:p>", styles=STYLES_EDGE,
                 core='<dcterms:created xsi:type="dcterms:W3CDTF"/>')
    # Styles without a default paragraph style: paragraph.style is None.
    docx_package("nodefault.docx", "<w:p>" + run_xml("<w:t>orphan</w:t>") + "</w:p>",
                 styles='<w:style w:type="paragraph" w:styleId="A"><w:name w:val="A"/></w:style>')
    # A default style without a name: style.name is None.
    docx_package("noname.docx", "<w:p>" + run_xml("<w:t>nameless</w:t>") + "</w:p>",
                 styles='<w:style w:type="paragraph" w:default="1" w:styleId="N"/>')
    # Table problems.
    docx_package("notblgrid.docx", "<w:tbl><w:tr><w:tc><w:p/></w:tc></w:tr></w:tbl>", styles=STYLES_EDGE)
    docx_package("vmergetop.docx",
                 "<w:tbl><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc></w:tr></w:tbl>",
                 styles=STYLES_EDGE)
    docx_package("badalign.docx", '<w:p><w:pPr><w:jc w:val="start"/></w:pPr>' + run_xml("<w:t>s</w:t>") + "</w:p>",
                 styles=STYLES_EDGE)
    docx_package("badbold.docx", "<w:p>" + run_xml("<w:t>s</w:t>", '<w:rPr><w:b w:val="yes"/></w:rPr>') + "</w:p>",
                 styles=STYLES_EDGE)
    # Not packages / wrong packages.
    with open(p("plain.docx"), "w") as f:
        f.write("dummy docx")
    write_zip("emptyzip.docx", {"readme.txt": "not an OPC package"})
    write_zip("badxml.docx", {
        "[Content_Types].xml": CT_DOCX,
        "_rels/.rels": PKG_RELS_DOC % "",
        "word/document.xml": "<w:document " + W_NS + "><w:body><w:p></w:body></w:document>",
    })
    write_zip("missingpart.docx", {"[Content_Types].xml": CT_DOCX, "_rels/.rels": PKG_RELS_DOC % ""})
    make_basic_docx()
    import shutil
    shutil.copy(p("basic.docx"), p("basic_copy.doc"))
    with open(p("notes.txt"), "w") as f:
        f.write("text")


def docx_cases():
    parser = DOCXParser()
    files = ["basic.docx", "edge.docx", "nostyles.docx", "dates.docx", "emptydate.docx", "nodefault.docx",
             "noname.docx", "notblgrid.docx", "vmergetop.docx", "badalign.docx", "badbold.docx", "plain.docx",
             "emptyzip.docx", "badxml.docx", "missingpart.docx", "basic_copy.doc", "notes.txt",
             "does_not_exist.docx", "basic.pptx", "data.xlsx"]
    for f in files:
        run("docx", f"parse:{f}", lambda f=f: parser.parse(p(f)), file=f, call="parse", options={})
    for f in ["basic.docx", "edge.docx", "nostyles.docx", "badalign.docx", "badbold.docx", "noname.docx"]:
        opts = {"extract_formatting": True}
        run("docx", f"parse_formatting:{f}", lambda f=f: parser.parse(p(f), **opts), file=f, call="parse",
            options=opts)
    for f in ["basic.docx", "notblgrid.docx", "vmergetop.docx"]:
        opts = {"extract_tables": False}
        run("docx", f"parse_no_tables:{f}", lambda f=f: parser.parse(p(f), **opts), file=f, call="parse",
            options=opts)
    run("docx", "extract_text:basic.docx", lambda: parser.extract_text(p("basic.docx")), file="basic.docx",
        call="extract_text", options={})
    run("docx", "extract_tables:basic.docx", lambda: parser.extract_tables(p("basic.docx")), file="basic.docx",
        call="extract_tables", options={})
    run("docx", "extract_tables:edge.docx", lambda: parser.extract_tables(p("edge.docx")), file="edge.docx",
        call="extract_tables", options={})
    run("docx", "extract_text:plain.docx", lambda: parser.extract_text(p("plain.docx")), file="plain.docx",
        call="extract_text", options={})


# ============================================================================ PPTX fixtures

PNG_1x1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360000002"
    "0000050001e2269b6b0000000049454e44ae426082"
)


def make_basic_pptx():
    prs = pptx.Presentation()
    cp = prs.core_properties
    cp.title = "Deck Title"
    cp.author = "Carol"
    cp.subject = "Planning"
    cp.keywords = "roadmap"
    cp.comments = "internal"
    cp.last_modified_by = "Dave"
    cp.created = dt.datetime(2023, 4, 5, 6, 7, 8)
    cp.modified = dt.datetime(2023, 5, 6, 7, 8, 9)
    # 1: title slide
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "  Welcome Deck  "
    s.placeholders[1].text = "Subtitle line\nSecond line"
    s.notes_slide.notes_text_frame.text = "Speaker notes one\nmore notes"
    # 2: title and content with bullets
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Agenda"
    body = s.placeholders[1].text_frame
    body.text = "Point A"
    para = body.add_paragraph()
    para.text = "Sub point"
    para.level = 1
    para = body.add_paragraph()
    para.add_run().text = "Run one "
    para.add_line_break()
    para.add_run().text = "after break"
    tb = s.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1))
    tb.text_frame.text = "Text box words"
    auto = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5), Inches(5), Inches(2), Inches(1))
    auto.text_frame.text = "Shape text"
    s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7), Inches(5), Inches(1), Inches(1))
    s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(1), Inches(2), Inches(2))
    # 3: blank layout, table, picture, chart, group, no title placeholder
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(4), Inches(1))
    tb.text_frame.text = "First box acts as title"
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "h1"
    tbl.cell(0, 1).text = "h2"
    tbl.cell(1, 0).text = "v1"
    pic = s.shapes.add_picture(io.BytesIO(PNG_1x1), Inches(6), Inches(1), Inches(1), Inches(1))
    chart_data = CategoryChartData()
    chart_data.categories = ["a", "b"]
    chart_data.add_series("S1", (1, 2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(3), Inches(2), chart_data)
    grp = s.shapes.add_group_shape()
    inner = grp.shapes.add_textbox(Inches(5), Inches(4), Inches(2), Inches(1))
    inner.text_frame.text = "inside group"
    s.notes_slide.notes_text_frame.text = ""
    # 4: picture placeholder filled through insert_picture (inherited size)
    s = prs.slides.add_slide(prs.slide_layouts[8])
    s.shapes.title.text = "Picture slide"
    s.placeholders[1].insert_picture(io.BytesIO(PNG_1x1))
    s.placeholders[2].text = "caption"
    # 5: empty title placeholder, text elsewhere
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = ""
    tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    tb.text_frame.text = "  only text  "
    prs.save(p("basic.pptx"))


P_NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'


def patch_pptx(src, dst, replacements):
    """Copy a package, applying (member, old, new) text replacements or adding members."""
    with zipfile.ZipFile(p(src)) as zin, zipfile.ZipFile(p(dst), "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        for name in names:
            data = zin.read(name)
            for member, old, new in replacements:
                if member == name and old is not None:
                    text = data.decode("utf-8")
                    assert old in text, (member, old)
                    data = text.replace(old, new, 1).encode("utf-8")
            zout.writestr(name, data)
        for member, old, new in replacements:
            if old is None and member not in names:
                zout.writestr(member, new)


def make_pptx_edge_cases():
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    tb.text_frame.text = "edge text"
    s.shapes.add_picture(io.BytesIO(PNG_1x1), Inches(5), Inches(1), Inches(1), Inches(1))
    s.notes_slide.notes_text_frame.text = "edge notes"
    prs.save(p("edge_base.pptx"))
    sp_plain = (
        '<p:sp><p:nvSpPr><p:cNvPr id="90" name="Plain"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/>'
        '<p:txBody><a:bodyPr/><a:p><a:r><a:t>no geometry</a:t></a:r><a:fld id="{1}" type="slidenum"><a:t>7</a:t></a:fld>'
        "<a:br/><a:r><a:t/></a:r></a:p></p:txBody></p:sp>"
    )
    patch_pptx("edge_base.pptx", "edge_shape.pptx", [("ppt/slides/slide1.xml", "</p:spTree>", sp_plain + "</p:spTree>")])
    # Notes slide whose body placeholder was removed.
    with zipfile.ZipFile(p("edge_base.pptx")) as z:
        notes = z.read("ppt/notesSlides/notesSlide1.xml").decode("utf-8")
    notes_no_body = re.sub(r'type="body"', 'type="dt"', notes)
    patch_pptx("edge_base.pptx", "edge_notes.pptx", [("ppt/notesSlides/notesSlide1.xml", notes, notes_no_body)])
    # Picture whose blip has no embed id.
    with zipfile.ZipFile(p("edge_base.pptx")) as z:
        slide = z.read("ppt/slides/slide1.xml").decode("utf-8")
    m = re.search(r'r:embed="rId\d+"', slide)
    patch_pptx("edge_base.pptx", "edge_noblip.pptx", [("ppt/slides/slide1.xml", m.group(0), "")])
    # Picture pointing to a WEBP image (unsupported by python-pptx's Image.ext).
    rels = None
    with zipfile.ZipFile(p("edge_base.pptx")) as z:
        rels = z.read("ppt/slides/_rels/slide1.xml.rels").decode("utf-8")
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
    webp = b"RIFF\x1a\x00\x00\x00WEBPVP8L\x0d\x00\x00\x00\x2f\x00\x00\x00\x10\x07\x10\x11\x11\x88\x88\xfe\x07\x00"
    patch_pptx("edge_base.pptx", "edge_webp.pptx", [(media[0], None, None)])
    with zipfile.ZipFile(p("edge_webp.pptx")) as z:
        members = {n: z.read(n) for n in z.namelist()}
    members[media[0]] = webp
    with zipfile.ZipFile(p("edge_webp.pptx"), "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in members.items():
            z.writestr(n, d)
    # A placeholder with an idx Python cannot convert.
    sp_bad_idx = (
        '<p:sp><p:nvSpPr><p:cNvPr id="91" name="BadIdx"/><p:cNvSpPr/><p:nvPr><p:ph type="body" idx="x1"/></p:nvPr></p:nvSpPr>'
        '<p:spPr/><p:txBody><a:bodyPr/><a:p><a:r><a:t>bad idx</a:t></a:r></a:p></p:txBody></p:sp>'
    )
    patch_pptx("edge_base.pptx", "edge_badidx.pptx", [("ppt/slides/slide1.xml", "</p:spTree>", sp_bad_idx + "</p:spTree>")])
    with open(p("plain.pptx"), "w") as f:
        f.write("not a presentation")


def pptx_cases():
    parser = PPTXParser()
    files = ["basic.pptx", "edge_base.pptx", "edge_shape.pptx", "edge_notes.pptx", "edge_noblip.pptx",
             "edge_webp.pptx", "edge_badidx.pptx", "plain.pptx", "basic.docx", "missing.pptx", "notes.txt"]
    all_opts = {"extract_images": True, "extract_shapes": True, "extract_notes": True}
    for f in files:
        run("pptx", f"parse:{f}", lambda f=f: parser.parse(p(f)), file=f, call="parse", options={})
        run("pptx", f"parse_all:{f}", lambda f=f: parser.parse(p(f), **all_opts), file=f, call="parse",
            options=all_opts)
    no_notes = {"extract_notes": False}
    run("pptx", "parse_no_notes:edge_notes.pptx", lambda: parser.parse(p("edge_notes.pptx"), **no_notes),
        file="edge_notes.pptx", call="parse", options=no_notes)
    run("pptx", "extract_text:basic.pptx", lambda: parser.extract_text(p("basic.pptx")), file="basic.pptx",
        call="extract_text", options={})
    shapes_only = {"extract_shapes": True}
    run("pptx", "extract_slides:basic.pptx", lambda: parser.extract_slides(p("basic.pptx"), **shapes_only),
        file="basic.pptx", call="extract_slides", options=shapes_only)


# ============================================================================ Excel fixtures


def styled_number_formats(wb):
    from openpyxl.styles import NamedStyle  # noqa: F401
    return wb


def make_data_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["name", "age", "score", "active", "joined", "notes", "code"])
    ws.append(["Alice", 30, 88.5, True, dt.datetime(2024, 1, 15), "NA", "007"])
    ws.append(["Bob", None, 92, False, dt.datetime(2024, 2, 1, 13, 30), "  spaced  ", "12"])
    ws.append(["Carol", 41, None, None, None, "x", "13"])
    ws.append([None, None, None, None, None, None, None])
    ws.append(["Dave", "29", 70.25, "TRUE", dt.date(2023, 12, 31), "", "14"])
    ws.append([None, None, None, None, None, None, None])
    ws["H9"].number_format = "0.00"  # a styled empty cell far away
    nums = wb.create_sheet("Numbers")
    nums.append([2020, 2021.5, "dup", "dup", None, "dup", True])
    nums.append([1, 2, 3, 4, 5, 6, 7])
    nums.append([1.5, None, "a", "b", "c", "d", "e"])
    nums.append([dt.time(10, 30), dt.time(11, 0, 5), 1, 2, 3, 4, 5])
    nums["A6"] = "after gap"
    wb.create_sheet("Empty")
    merged = wb.create_sheet("Merged")
    merged.append(["h1", "h2", "h3"])
    merged.append(["a", "b", "c"])
    merged.append(["d", "e", "f"])
    merged.merge_cells("A2:B3")
    merged["E7"].number_format = "0%"
    ragged = wb.create_sheet("Ragged")
    ragged.append(["a", "b", "c"])
    ragged.append([1])
    ragged.append([1, 2, 3, 4])
    ragged.append(["only first"])
    one = wb.create_sheet("OneCol")
    one.append(["v"])
    one.append([1])
    one.append([None])
    one.append(["  "])
    one.append([2])
    mixed = wb.create_sheet("Mixed")
    mixed.append(["objs", "bools", "boolstr", "dates", "na", "big", "floatint"])
    mixed.append([1, True, "True", dt.datetime(2024, 1, 1), "N/A", 10 ** 15, 1.0])
    mixed.append(["x", False, "false", "text", "null", -5, 2.5])
    mixed.append([2.5, True, "NA", dt.datetime(2024, 1, 3), "#N/A", 7, None])
    mixed.append([True, None, "TRUE", None, "", 0, 3])
    wb.properties.creator = "fixture"
    wb.save(p("data.xlsx"))


SHEET_NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'


def make_crafted_xlsx():
    """A workbook with hand-written cells: formulas with cached values, error/bool/inline/ISO
    cells, rich shared strings, number formats (dates, times, timedeltas, numFmt 27) and rows
    or cells without coordinates."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw"
    ws.append(["placeholder"])
    wb.save(p("crafted_base.xlsx"))
    with zipfile.ZipFile(p("crafted_base.xlsx")) as z:
        members = {n: z.read(n) for n in z.namelist()}
    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<styleSheet {SHEET_NS}>'
        '<numFmts count="4"><numFmt numFmtId="164" formatCode="yyyy-mm-dd"/>'
        '<numFmt numFmtId="165" formatCode="[h]:mm:ss"/>'
        '<numFmt numFmtId="166" formatCode="&quot;day&quot; 0.00"/>'
        '<numFmt numFmtId="14" formatCode="0.000"/></numFmts>'
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="8">'
        '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        '<xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
        '<xf numFmtId="20" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
        '<xf numFmtId="165" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
        '<xf numFmtId="166" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
        '<xf numFmtId="27" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
        '<xf numFmtId="14" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
        '<xf numFmtId="22" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
        '</cellXfs>'
        '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
        '</styleSheet>'
    )
    shared = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<sst {SHEET_NS} count="5" uniqueCount="5">'
        '<si><t>label</t></si>'
        '<si><r><rPr><b/></rPr><t>rich </t></r><r><t>text</t></r></si>'
        '<si><t>esc_x005F_x000D_ape</t></si>'
        '<si><t xml:space="preserve">  padded </t><rPh sb="0" eb="1"><t>phonetic</t></rPh></si>'
        '<si><t/></si>'
        '</sst>'
    )
    sheet = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<worksheet {SHEET_NS}><sheetData>'
        '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>2</v></c>'
        '<c r="D1" t="s"><v>3</v></c><c r="E1" t="inlineStr"><is><t>inline</t></is></c><c r="F1" t="s"><v>4</v></c>'
        '<c r="G1" t="str"><f>"x"&amp;"y"</f><v>xy</v></c></row>'
        '<row r="2"><c r="A2" s="1"><v>45306</v></c><c r="B2" s="2"><v>0.5</v></c><c r="C2" s="3"><v>1.25</v></c>'
        '<c r="D2" s="4"><v>3.5</v></c><c r="E2" s="5"><v>45306</v></c><c r="F2" s="6"><v>45306</v></c>'
        '<c r="G2" s="7"><v>45306.75</v></c></row>'
        '<row r="3"><c r="A3" t="e"><v>#DIV/0!</v></c><c r="B3" t="b"><v>1</v></c><c r="C3" t="b"><v>0</v></c>'
        '<c r="D3"><f>1+1</f><v>2</v></c><c r="E3"><f>A9</f></c><c r="F3" t="d"><v>2024-03-04T05:06:07</v></c>'
        '<c r="G3" t="d"><v>2024-03-05</v></c></row>'
        '<row><c><v>1.0</v></c><c><v>2E2</v></c><c t="e"><v>#N/A</v></c><c r="F4"><v>-7</v></c></row>'
        '<row r="6" spans="1:3" ht="20" customHeight="1"/>'
        '<row r="7"><c r="B7" s="1"><v>59</v></c><c r="C7" s="1"><v>61</v></c><c r="D7" s="1"><v>0</v></c></row>'
        '<row r="8"><c r="A8" s="3"><v>2.000001</v></c><c r="B8" t="d"><v>10:20:30</v></c><c r="C8" t="d"><v>PT1H30M</v></c></row>'
        '</sheetData></worksheet>'
    )
    members["xl/styles.xml"] = styles.encode()
    members["xl/sharedStrings.xml"] = shared.encode()
    members["xl/worksheets/sheet1.xml"] = sheet.encode()
    ct = members["[Content_Types].xml"].decode()
    if "sharedStrings" not in ct:
        ct = ct.replace("</Types>",
                        '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
    members["[Content_Types].xml"] = ct.encode()
    rels = members["xl/_rels/workbook.xml.rels"].decode()
    if "sharedStrings" not in rels:
        rels = rels.replace("</Relationships>",
                            '<Relationship Id="rIdSS" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
    members["xl/_rels/workbook.xml.rels"] = rels.encode()
    with zipfile.ZipFile(p("crafted.xlsx"), "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in members.items():
            z.writestr(n, d)
    # The same workbook with the 1904 date system.
    wbxml = members["xl/workbook.xml"].decode()
    if re.search(r"<workbookPr[ />]", wbxml):
        wbxml = re.sub(r"<workbookPr( [^>]*)?/>", lambda m: '<workbookPr%s date1904="1"/>' % (m.group(1) or ""), wbxml)
    else:
        wbxml = wbxml.replace("<sheets>", '<workbookPr date1904="1"/><sheets>')
    members1904 = dict(members)
    members1904["xl/workbook.xml"] = wbxml.encode()
    with zipfile.ZipFile(p("crafted1904.xlsx"), "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in members1904.items():
            z.writestr(n, d)
    os.remove(p("crafted_base.xlsx"))


def make_excel_fixtures():
    import shutil
    make_data_xlsx()
    make_crafted_xlsx()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["k", "v"])
    ws.append(["a", 1])
    cs = wb.create_chartsheet("Chart")
    from openpyxl.chart import BarChart, Reference
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    cs.add_chart(chart)
    wb.save(p("chart.xlsx"))
    shutil.copy(p("data.xlsx"), p("data_copy.xlsm"))
    shutil.copy(p("data.xlsx"), p("data_copy.xls"))
    with open(p("ole.xls"), "wb") as f:
        f.write(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 504)
    with open(p("text.xlsx"), "w") as f:
        f.write("not a workbook")
    with open(p("empty.xlsx"), "w") as f:
        pass
    write_zip("plainzip.xlsx", {"a.txt": "x"})
    write_zip("nocontenttypes.xlsx", {"xl/workbook.xml": "<workbook/>"})


def excel_cases():
    parser = ExcelParser()

    def case(name, file, sheet=None, **opts):
        def call():
            if sheet is None:
                return parser.parse(p(file), **opts)
            return parser.parse(p(file), sheet_name=sheet, **opts)
        run("excel", name, call, file=file, call="parse", sheet_name=sheet, options=opts)

    for f in ["data.xlsx", "crafted.xlsx", "crafted1904.xlsx", "chart.xlsx", "data_copy.xlsm", "data_copy.xls",
              "ole.xls", "text.xlsx", "empty.xlsx", "plainzip.xlsx", "nocontenttypes.xlsx", "missing.xlsx",
              "notes.txt"]:
        case(f"pandas:{f}", f)
        case(f"openpyxl:{f}", f, engine="openpyxl")
    for sheet in ["People", "Numbers", "Empty", "Merged", "Ragged", "OneCol", "Mixed", "Nope"]:
        case(f"pandas:data.xlsx:{sheet}", "data.xlsx", sheet)
        case(f"openpyxl:data.xlsx:{sheet}", "data.xlsx", sheet, engine="openpyxl")
    case("pandas:chart.xlsx:Chart", "chart.xlsx", "Chart")
    case("openpyxl:chart.xlsx:Chart", "chart.xlsx", "Chart", engine="openpyxl")
    case("openpyxl:chart.xlsx:Data", "chart.xlsx", "Data", engine="openpyxl")
    variants = [
        {"header_row": 1},
        {"header_row": 2, "skip_rows": 1},
        {"skip_rows": 2},
        {"skip_rows": 50},
        {"header_row": 50},
        {"max_rows": 1},
        {"max_rows": 2, "skip_rows": 1},
        {"max_rows": 0},
        {"max_rows": -1},
        {"header_row": None},
        {"header_row": None, "max_rows": 2},
        {"header_row": -1},
        {"header_row": True},
        {"header_row": "x"},
        {"skip_rows": None},
    ]
    for v in variants:
        key = ",".join(f"{k}={v[k]}" for k in v)
        case(f"pandas:data.xlsx:People:{key}", "data.xlsx", "People", **v)
        case(f"pandas:data.xlsx:Numbers:{key}", "data.xlsx", "Numbers", **v)
        if not any(isinstance(x, str) for x in v.values()):
            case(f"openpyxl:data.xlsx:People:{key}", "data.xlsx", "People", engine="openpyxl", **v)
    case("openpyxl:data.xlsx:Merged:max_rows=1", "data.xlsx", "Merged", engine="openpyxl", max_rows=1)
    case("pandas:crafted.xlsx:Raw:max_rows=3", "crafted.xlsx", "Raw", max_rows=3)
    case("engine_other:data.xlsx", "data.xlsx", "People", engine="calamine")
    run("excel", "extract_sheet:data.xlsx:Numbers", lambda: parser.extract_sheet(p("data.xlsx"), "Numbers"),
        file="data.xlsx", call="extract_sheet", sheet_name="Numbers", options={})
    run("excel", "extract_sheet:data.xlsx:empty", lambda: parser.extract_sheet(p("data.xlsx"), ""),
        file="data.xlsx", call="extract_sheet", sheet_name="", options={})
    run("excel", "extract_sheet:data.xlsx:missing", lambda: parser.extract_sheet(p("data.xlsx"), "Nope"),
        file="data.xlsx", call="extract_sheet", sheet_name="Nope", options={})


def main():
    make_docx_edge_cases()
    make_basic_pptx()
    make_pptx_edge_cases()
    make_excel_fixtures()
    docx_cases()
    pptx_cases()
    excel_cases()
    with open(os.path.join(HERE, "oracle.json"), "w") as f:
        json.dump(OUT, f, indent=1, sort_keys=False, ensure_ascii=False)
        f.write("\n")


if __name__ == "__main__":
    main()
